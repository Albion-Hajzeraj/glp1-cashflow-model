"""Forecast, levers, scenarios and doubling bridge for omzo_model.xlsx (live formulas).

Engine (one sheet per run, identical layout):
- Rows = segments: signup cohort x channel x plan. Existing cohorts start from actual
  active customers in the last actual month (Data_ActiveEnd). New cohorts come from
  spend / CPA by channel, split by plan mix.
- Each month an active segment keeps (1 - churn), where churn = 1 - r(k)/r(k-1) from the
  fitted retention curve for its channel x plan (Data_RetentionPlan), adjusted by the
  levers (support check-ins, dunning, pause, price response). Paying customers ship
  monthly; monthly-plan customers are charged monthly, prepaid every `term` months.
- Revenue, variable costs, cash timing as in the actuals: card payouts T+n, pharmacy
  paid the next month (net-30), ads paid with the ad lag, opex same month.

Runs: Forecast (scenario switch on Inputs), Eng_Base, Eng_Down (fixed scenarios for the
side-by-side table), Eng_Bridge0..7 (levers switched on cumulatively in a fixed order;
Bridge(8) = Forecast). The bridge bars therefore telescope exactly to the total change.
"""
from __future__ import annotations

from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font

NF = 9  # forecast months (Oct 2026 - Jun 2027)

LEVERS = [
    # key, name, main value, unit, start, ramp, confidence, justification, evidence
    ("a", "Channel reallocation", 1.0, "share of planned cuts executed", "2026-10", 2, "high",
     "TikTok CPA is above its 4-month max CPA and ~2x target (slowest payback); meta is 54% over target. "
     "Shift budget to Google, affiliate and referral, where CPA is closest to target.",
     "Geo/holdout incrementality tests; marginal CPA curves by channel; Google impression-share headroom; "
     "affiliate partner capacity."),
    ("b1", "Retention: titration & side-effect check-ins", 0.15, "relative churn cut, tenure months 1-3",
     "2026-11", 3, "medium",
     "Churn is highest in months 1-3, when doses escalate and GI side effects peak; side effects are "
     "~25% of voluntary cancels. Proactive clinician check-ins address the cause.",
     "A/B test on new cohorts; cancel-reason and side-effect tickets; month-1-3 churn by arm."),
    ("b2", "Retention: dunning (failed-payment recovery)", 0.15, "pp added to recovery rate", "2026-10", 2, "high",
     "40% of failed renewals are lost today (60% recovery). Card updater, smart retry timing and SMS "
     "reminders typically lift recovery 10-20pp.",
     "Decline-code mix; recovery rate by retry day; processor account-updater coverage."),
    ("b3", "Retention: pause instead of cancel", 0.20, "share of would-be cancels offered a pause who take it",
     "2026-11", 2, "medium",
     "Some cancels are temporary (travel, cost, side-effect break). A clinically appropriate 1-month "
     "pause keeps the relationship; 60% of pausers are assumed to resume.",
     "Cancel-flow test: pause take-up and resume rate; outcomes of paused patients."),
    ("c", "Plan mix: shift to 3-month prepaid", 0.15, "pp added to prepaid share of new customers",
     "2026-11", 3, "medium",
     "Prepaid collects 3 months up front and retains better through month 3. An extra 5% off "
     "($852 vs $897) is assumed to move 15pp of new customers.",
     "Checkout price test (take rate vs discount); prepaid vs monthly retention at equal tenure."),
    ("d", "Pricing: modest list-price increase", 0.05, "price change (both plans)", "2027-01", 1, "low",
     "$349 sits at or below typical compounded GLP-1 programs. +5% is modest. Assumed response: "
     "4% fewer sign-ups (elasticity -0.8) and 2.5% higher churn.",
     "Price test on new visitors; competitor price scan; churn of a grandfathered vs new-price cohort."),
    ("e", "Medication cost renegotiation", 0.08, "cut in medication base cost", "2027-01", 1, "medium",
     "Volume is up ~6x since January. A second-source pharmacy or volume tier at the next contract "
     "review is assumed to cut 8% off the $98 base.",
     "Written RFQ quotes from 2+ 503B/503A pharmacies; quality and compliance audit; contract terms."),
    ("f", "Shipping & fulfillment savings", 0.12, "cut in shipping + packaging per shipment", "2026-12", 2, "medium",
     "Regional cold-chain carriers and right-sized insulated packaging cut cost per shipment.",
     "Carrier quotes by zone; cold-chain validation (temperature logs); damage/reship rates."),
]
LEVER_KEYS = [lv[0] for lv in LEVERS]
# Plan-mix lever is modelled but OFF by default: at the target month it is cash-negative
# (prepaid earns $284-299/month vs $349 and the retention gain does not offset it).
DEFAULT_OFF = {"c"}
SCEN_PARAMS = [("Lever effect multiplier", 1.0, 0.5, "x of planned lever effect"),
               ("Medication cost shock", 0.0, 0.15, "added to medication base cost"),
               ("Retention drop", 0.0, 0.10, "points off the retention curve, months 1+"),
               ("CPA shock", 0.0, 0.20, "added to baseline CPA")]


def L(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)


def fcol(f):
    """Forecast month f (1..NF) column on engine/Levers month grids: D..L."""
    return L(3 + f)


class Forecast:
    def __init__(self, m):
        self.m, self.wb, self.R, self.D = m, m.wb, m.R, m.D
        self.nm, self.channels = m.nm, m.channels

    # ------------------------------------------------------------------ Levers sheet
    def levers(self):
        ws, R, D, m = self.wb["Levers"], self.R, self.D, self.m
        ws["A1"] = "Levers: plan to double monthly net operating cash flow (blue = input)"
        ws["A1"].font = m.TITLE
        ws["A2"] = ("Guardrail: retention and pricing actions must respect clinical suitability. No incentives to stay on "
                    "medication against clinical advice, no bypassing eligibility screening; pause/check-ins are "
                    "clinician-led and patients can always stop treatment.")
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[2].height = 32
        r = 4
        m.section(ws, r, "Forecast months", 12); r += 1
        m.header(ws, r, ["Forecast month #"] + list(range(1, NF + 1)), start_col=3)
        self.f_idx_row = r
        r += 1
        ws.cell(r, 1, "Forecast month")
        cal = D["Data_Calendar"]["month"]
        for f in range(1, NF + 1):
            ws.cell(r, 3 + f, f"=INDEX({cal},{R['n_months']}+{fcol(f)}{self.f_idx_row})")
        self.fm_row = r
        self.fm_rng = f"Levers!$D${r}:${fcol(NF)}${r}"
        r += 2

        m.section(ws, r, "Forecast baseline (black = formula from actuals; overwrite to change)", 12); r += 1
        def scalar(key, label, value, unit, fmt=None, is_input=True):
            nonlocal r
            ws.cell(r, 1, label)
            c = ws.cell(r, 2, value)
            if is_input:
                c.font = m.BLUE
            if fmt:
                c.number_format = fmt
            ws.cell(r, 3, unit)
            R[key] = f"Levers!$B${r}"
            r += 1
        scalar("fc_window", "Baseline CPA window", 3, "last N actual months")
        scalar("fc_opex_g", "Fixed opex growth", 0.0, "per month", "0.0%")
        scalar("fc_lambda", "Share of a month's charges paid out next month", f"={R['payout_lag']}/{R['days_per_month']}",
               "T+n / days per month", "0.0%", False)
        scalar("fc_carry_card", "Card payouts from actual charges landing in month 1",
               f"=SUMIFS({D['Data_Payments']['amount']},{D['Data_Payments']['payout_idx']},\">\"&{R['n_months']})",
               "$", m.USD, False)
        scalar("fc_carry_pharm", "Pharmacy payable at end of actuals (paid in month 1)",
               f"=SUMIFS({D['Data_Shipments']['pharmacy']},{D['Data_Shipments']['paid_idx']},\">\"&{R['n_months']})",
               "$", m.USD, False)
        scalar("fc_titr_window", "Check-in lever applies to tenure months up to", 3, "months")
        scalar("fc_pause_return", "Pausers who resume", 0.60, "share", "0.0%")
        scalar("fc_c_disc", "Extra prepaid discount (plan-mix lever)", 0.05, "off prepaid price", "0.0%")
        scalar("fc_d_demand", "Price elasticity of sign-ups", -0.8, "% sign-ups per % price")
        scalar("fc_d_churn", "Churn response to price", 0.5, "% churn per % price")
        r += 1
        ws.cell(r, 1, "Forecast month")
        for f in range(1, NF + 1):
            ws.cell(r, 3 + f, f"={fcol(f)}${self.fm_row}")
        r += 1
        ws.cell(r, 1, "Medication base cost ($ / supply)")
        for f in range(1, NF + 1):
            ws.cell(r, 3 + f, f"=INDEX({R['med_base']},1,{R['n_months']})").number_format = m.USD2
        self.med_row = r; r += 1
        ws.cell(r, 1, "Fixed opex ($ / month)")
        for f in range(1, NF + 1):
            ws.cell(r, 3 + f, (f"=SUM(INDEX(Inputs!$B${R['opex_r0']}:${L(1 + self.nm)}${R['opex_r1']},0,{R['n_months']}))"
                               f"*(1+{R['fc_opex_g']})^{fcol(f)}${self.f_idx_row}")).number_format = m.USD
        self.opex_row = r; r += 2

        m.header(ws, r, ["Channel", "Baseline spend / mo (last actual)", "Baseline CPA (window)", "Prepaid mix",
                         "Paid share", "Promo share (monthly)", "CPA elasticity to spend", "Cut (lever a)",
                         "Receiver weight (lever a)", "Spend freed at full cut"])
        r += 1
        self.ch0 = r
        Cu = D["Data_Customers"]
        cuts = {"tiktok": 0.50, "meta": 0.15}
        weights = {"google": 0.50, "affiliate": 0.30, "organic_referral": 0.20}
        elast = {"meta": 0.25, "google": 0.35, "tiktok": 0.25, "affiliate": 0.30, "organic_referral": 0.60}
        for ch in self.channels:
            sp_row = R["spend_r0"] + self.channels.index(ch)
            ws.cell(r, 1, ch)
            ws.cell(r, 2, f"=INDEX(Inputs!$B${sp_row}:${L(1 + self.nm)}${sp_row},1,{R['n_months']})").number_format = m.USD
            ws.cell(r, 3, (f"=SUMIFS(Inputs!$B${sp_row}:${L(1 + self.nm)}${sp_row},{R['month_idx_row']},\">=\"&({R['n_months']}"
                           f"-{R['fc_window']}+1))/COUNTIFS({Cu['channel']},$A{r},{Cu['signup_idx']},\">=\"&({R['n_months']}"
                           f"-{R['fc_window']}+1))")).number_format = m.USD2
            ls = m.LS[("base", ch)]
            ws.cell(r, 4, f"={ls['mix']}").number_format = "0.0%"
            ws.cell(r, 5, f"={ls['paid']}").number_format = "0.0%"
            ws.cell(r, 6, f"={ls['promo']}").number_format = "0.0%"
            for j, v in enumerate([elast[ch], cuts.get(ch, 0.0), weights.get(ch, 0.0)]):
                c = ws.cell(r, 7 + j, v)
                c.font, c.number_format = m.BLUE, ("0.00" if j == 0 else "0.0%")
            ws.cell(r, 10, f"=B{r}*H{r}").number_format = m.USD
            r += 1
        self.ch1 = r - 1
        ws.cell(r, 1, "Total").font = m.BOLD
        for col in "BJ":
            ws[f"{col}{r}"] = f"=SUM({col}{self.ch0}:{col}{self.ch1})"
            ws[f"{col}{r}"].number_format = m.USD
        ws[f"I{r}"] = f"=SUM(I{self.ch0}:I{self.ch1})"
        ws[f"I{r}"].number_format = "0.0%"
        self.ch_total = r
        R["ch_list_fc"] = f"Levers!$A${self.ch0}:$A${self.ch1}"
        r += 2

        m.section(ws, r, "Levers (each with value, start month and ramp-up)", 12); r += 1
        m.header(ws, r, ["Lever", "On (1/0)", "Value", "Unit", "Start month", "Ramp (months)", "Confidence",
                         "Justification", "Evidence needed to validate"])
        r += 1
        self.lv_rows = {}
        for key, name, val, unit, start, ramp, conf, why, ev in LEVERS:
            ws.cell(r, 1, f"{key}) {name}")
            on_default = 0 if key in DEFAULT_OFF else 1
            for col, v, fmt in [(2, on_default, None), (3, val, "0.0%"), (5, start, None), (6, ramp, None), (7, conf, None)]:
                c = ws.cell(r, col, v)
                c.font = m.BLUE
                if fmt:
                    c.number_format = fmt
            ws.cell(r, 4, unit)
            ws.cell(r, 8, why).alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(r, 9, ev).alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 62
            self.lv_rows[key] = r
            r += 1
        r += 1
        m.section(ws, r, "Ramp factor by month (0 before start, linear to 100% over the ramp)", 12); r += 1
        ws.cell(r, 1, "Lever")
        for f in range(1, NF + 1):
            ws.cell(r, 3 + f, f"={fcol(f)}${self.fm_row}").font = m.BOLD
        r += 1
        self.ramp_rows = {}
        for key, *_ in LEVERS:
            lr = self.lv_rows[key]
            ws.cell(r, 1, f"=A{lr}")
            for f in range(1, NF + 1):
                fc = fcol(f)
                ws.cell(r, 3 + f, (f"=IFERROR(IF({fc}${self.f_idx_row}<MATCH($E${lr},{self.fm_rng},0),0,"
                                   f"MIN(1,({fc}${self.f_idx_row}-MATCH($E${lr},{self.fm_rng},0)+1)/MAX(1,$F${lr}))),0)")
                                   ).number_format = "0%"
            self.ramp_rows[key] = r
            r += 1
        ws.column_dimensions["A"].width = 46
        for col, w in zip("BCDEFGHIJKL", [14, 14, 30, 12, 12, 12, 60, 60, 14, 14]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ Scenario params
    def scenario_params(self):
        ws, m = self.wb["Scenarios"], self.m
        ws["A1"] = "Scenarios: base vs downside (scenario switch on Inputs drives the Forecast sheet)"
        ws["A1"].font = m.TITLE
        m.header(ws, 4, ["Parameter", "Unit", "Base", "Downside"])
        self.sp_rows = []
        for i, (lab, base, down, unit) in enumerate(SCEN_PARAMS):
            r = 5 + i
            ws.cell(r, 1, lab)
            ws.cell(r, 2, unit)
            for col, v in [(3, base), (4, down)]:
                c = ws.cell(r, col, v)
                c.font, c.number_format = m.BLUE, "0.0%"
            self.sp_rows.append(r)

    # ------------------------------------------------------------------ engine
    def engine(self, name, scen_formula, flag_formulas, title):
        ws = self.wb.create_sheet(name)
        m, R, D = self.m, self.R, self.D
        n = R["n_months"]
        ws["A1"] = title
        ws["A1"].font = m.TITLE
        ws["A3"], ws["B3"] = "Scenario # (1 = base, 2 = downside)", scen_formula
        for i, r in enumerate(self.sp_rows):
            ws.cell(4 + i, 1, SCEN_PARAMS[i][0])
            ws.cell(4 + i, 2, f"=INDEX(Scenarios!$C${r}:$D${r},1,$B$3)").number_format = "0.0%"
        EFF, MED, RET, CPA = "$B$4", "$B$5", "$B$6", "$B$7"
        m.header(ws, 9, ["Lever on-factor", "Flag", ""] + [f"=Levers!{fcol(f)}${self.fm_row}" for f in range(1, NF + 1)])
        on = {}
        for j, key in enumerate(LEVER_KEYS):
            r = 10 + j
            ws.cell(r, 1, f"=Levers!A{self.lv_rows[key]}")
            ws.cell(r, 2, flag_formulas[j])
            for f in range(1, NF + 1):
                ws.cell(r, 3 + f, f"=$B{r}*Levers!$B${self.lv_rows[key]}*{EFF}*Levers!{fcol(f)}${self.ramp_rows[key]}"
                        ).number_format = "0%"
            on[key] = r
        lv = lambda key, col="C": f"Levers!${col}${self.lv_rows[key]}"
        # month header + drivers
        m.header(ws, 19, ["Month driver", "", ""] + [f"=Levers!{fcol(f)}${self.fm_row}" for f in range(1, NF + 1)])
        ws.cell(20, 1, "Forecast month #")
        ws.cell(21, 1, "Absolute month #")
        for f in range(1, NF + 1):
            ws.cell(20, 3 + f, f)
            ws.cell(21, 3 + f, f"={n}+{fcol(f)}20")
        drivers = [
            ("price_m", "Monthly price", lambda c: f"{R['price_m']}*(1+{lv('d')}*{c}{on['d']})", m.USD2),
            ("price_p", "Prepaid list price", lambda c: f"{R['price_p']}*(1+{lv('d')}*{c}{on['d']})", m.USD2),
            ("pdisc", "Extra prepaid discount (new forecast cohorts only)", lambda c: f"{R['fc_c_disc']}*{c}{on['c']}",
             "0.0%"),
            ("med", "Medication base cost", lambda c: (f"Levers!{c}${self.med_row}*(1+{MED})*(1-{lv('e')}*{c}{on['e']})"),
             m.USD2),
            ("ship", "Shipping per shipment", lambda c: f"{R['shipping']}*(1-{lv('f')}*{c}{on['f']})", m.USD2),
            ("pack", "Packaging per shipment", lambda c: f"{R['packaging']}*(1-{lv('f')}*{c}{on['f']})", m.USD2),
            ("titr", "Check-in churn cut (months 1-window)", lambda c: f"{lv('b1')}*{c}{on['b1']}", "0.0%"),
            ("dun", "Dunning churn cut (renewal months)", lambda c: f"{R['fail_rate']}*{lv('b2')}*{c}{on['b2']}", "0.00%"),
            ("pause", "Share of churners pausing and resuming", lambda c: f"{lv('b3')}*{R['fc_pause_return']}*{c}{on['b3']}",
             "0.0%"),
            ("chp", "Churn multiplier from price", lambda c: f"1+{lv('d')}*{c}{on['d']}*{R['fc_d_churn']}", "0.000"),
            ("demand", "Sign-up multiplier from price", lambda c: f"MAX(0,1+{lv('d')}*{c}{on['d']}*{R['fc_d_demand']})",
             "0.000"),
            ("shift", "Prepaid mix shift", lambda c: f"{lv('c')}*{c}{on['c']}", "0.0%"),
            ("freed", "Spend freed by reallocation", lambda c: f"Levers!$J${self.ch_total}*{lv('a')}*{c}{on['a']}", m.USD),
        ]
        dr = {}
        r = 22
        for key, label, fn, fmt in drivers:
            ws.cell(r, 1, label)
            for f in range(1, NF + 1):
                ws.cell(r, 3 + f, "=" + fn(fcol(f))).number_format = fmt
            dr[key] = r
            r += 1
        # channel blocks
        ch_rows = {}
        r += 1
        nch = len(self.channels)
        blocks = [
            ("spend", "Spend by channel", m.USD),
            ("cpa", "CPA by channel (diminishing returns + scenario shock)", m.USD2),
            ("new", "New paying customers by channel", "#,##0.0"),
            ("mixp", "Prepaid share of new customers", "0.0%"),
            ("promo1", "Promo first charges (monthly plan)", "#,##0.0"),
        ]
        for key, label, fmt in blocks:
            ws.cell(r, 1, label).font = m.BOLD
            r += 1
            ch_rows[key] = r
            for i, ch in enumerate(self.channels):
                lr = self.ch0 + i
                ws.cell(r + i, 1, ch)
                for f in range(1, NF + 1):
                    c = fcol(f)
                    sp, cp, nw, mx = (f"{c}{ch_rows['spend'] + i}", f"{c}{ch_rows.get('cpa', 0) + i}",
                                      f"{c}{ch_rows.get('new', 0) + i}", f"{c}{ch_rows.get('mixp', 0) + i}")
                    fx = {
                        "spend": f"=Levers!$B${lr}*(1-Levers!$H${lr}*{lv('a')}*{c}${on['a']})+{c}${dr['freed']}*Levers!$I${lr}",
                        "cpa": (f"=Levers!$C${lr}*(1+{CPA})*IF(Levers!$B${lr}>0,({sp}/Levers!$B${lr})^Levers!$G${lr},1)"),
                        "new": f"=IF({cp}>0,{sp}/{cp},0)*{c}${dr['demand']}",
                        "mixp": f"=MIN(1,MAX(0,Levers!$D${lr}+{c}${dr['shift']}))",
                        "promo1": f"={nw}*(1-{mx})*Levers!$E${lr}*Levers!$F${lr}",
                    }[key]
                    ws.cell(r + i, 3 + f, fx).number_format = fmt
            r += nch + 1
        # retention curves for this scenario
        K = self.nm + NF            # k = 0..K-1
        rp = D["Data_RetentionPlan"]
        nseg = len(self.channels) * 2
        r += 1
        ws.cell(r, 1, "Retention curve r(k) for this scenario (channel x plan)").font = m.BOLD
        m.header(ws, r + 1, ["Channel", "Plan", "Key"] + list(range(K)))
        r0c = r + 2
        for i in range(nseg):
            rr = r0c + i
            ws.cell(rr, 1, f"=INDEX({rp['channel']},{i + 1})")
            ws.cell(rr, 2, f"=INDEX({rp['plan_type']},{i + 1})")
            ws.cell(rr, 3, f"=A{rr}&\"|\"&B{rr}")
            for k in range(K):
                src = f"INDEX(Data_RetentionPlan!$D$2:${L(3 + K)}${1 + nseg},{i + 1},{k + 1})"
                ws.cell(rr, 4 + k, "=1" if k == 0 else f"=MAX(0,{src}-{RET})").number_format = "0.000"
        r1c = r0c + nseg - 1
        key_rng = f"$C${r0c}:$C${r1c}"
        r = r1c + 2
        ws.cell(r, 1, "Monthly survival s(k) = r(k) / r(k-1)").font = m.BOLD
        s0 = r + 1
        for i in range(nseg):
            rr = s0 + i
            ws.cell(rr, 1, f"=A{r0c + i}")
            ws.cell(rr, 2, f"=B{r0c + i}")
            for k in range(K):
                cur, prev = f"{L(4 + k)}{r0c + i}", f"{L(3 + k)}{r0c + i}"
                ws.cell(rr, 4 + k, "=1" if k == 0 else f"=IF({prev}=0,0,{cur}/{prev})").number_format = "0.000"
        s_tbl = f"$D${s0}:${L(3 + K)}${s0 + nseg - 1}"
        r = s0 + nseg + 1

        # summary rows (filled after the segment table position is known)
        sum_start = r + 1
        names = ["new_total", "active", "paused", "shipments", "ship_k0", "charges_m", "charges_p", "first_m",
                 "first_p", "charges_p_new", "promo_first", "ship_prepaid", "ship_p_new", "dose_ship", "billing_gross",
                 "disc_promo", "disc_prepaid", "discount", "billing_net",
                 "first_net", "renew_net", "refunds", "chargebacks", "recognized", "net_revenue", "med", "disp",
                 "consult", "shipping", "packaging", "processing", "cbfees", "tvc", "contribution", "margin",
                 "marketing", "opex", "op_profit", "card", "cash_refunds", "receipts", "pharm_paid", "mkt_paid",
                 "other_paid", "opex_paid", "payments", "ncf", "opening", "closing"]
        S = {nm: sum_start + i for i, nm in enumerate(names)}
        seg0 = S["closing"] + 4
        segs = [(c, ch, p) for c in range(1, self.nm + NF + 1) for ch in self.channels
                for p in ["monthly", "prepaid_3m"]]
        seg1 = seg0 + len(segs) - 1
        blk = {b: 8 + i * NF for i, b in enumerate(["H", "A", "P", "S", "C", "D"])}  # first column index
        bc = lambda b, f: L(blk[b] + f - 1)

        # segment table
        hdr = ["Cohort #", "Channel", "Plan", "Key", "Curve row", "Active at start", "Paid share"]
        labels = {"H": "Churn", "A": "Active", "P": "Paused (lever)", "S": "Paid shipments", "C": "Charges",
                  "D": "Dose-weighted shipments"}
        m.header(ws, seg0 - 1, hdr + [f"{labels[b]} m{f}" for b in blk for f in range(1, NF + 1)])
        ae = D["Data_ActiveEnd"]
        term, tw = R["term_p"], R["fc_titr_window"]
        for i, (coh, ch, plan) in enumerate(segs):
            rr = seg0 + i
            ws.cell(rr, 1, coh)
            ws.cell(rr, 2, ch)
            ws.cell(rr, 3, plan)
            ws.cell(rr, 4, f"=B{rr}&\"|\"&C{rr}")
            ws.cell(rr, 5, f"=MATCH(D{rr},{key_rng},0)")
            ws.cell(rr, 6, (f"=SUMIFS({ae['active_customers']},{ae['cohort_idx']},A{rr},{ae['channel']},B{rr},"
                            f"{ae['plan_type']},C{rr})") if coh <= self.nm else 0)
            ws.cell(rr, 7, f"=INDEX(Levers!$E${self.ch0}:$E${self.ch1},MATCH(B{rr},{R['ch_list_fc']},0))")
            chm = f"MATCH($B{rr},{R['ch_list_fc']},0)"
            for f in range(1, NF + 1):
                c = fcol(f)
                k = f"({c}$21-$A{rr})"
                flag = f"IF($C{rr}=\"monthly\",1,IF({k}={term}*ROUND({k}/{term},0),1,0))"
                prev_a = f"F{rr}" if f == 1 else f"{bc('A', f - 1)}{rr}"
                h, a, p, s = (f"{bc(b, f)}{rr}" for b in "HAPS")
                new = (f"INDEX(${fcol(1)}${ch_rows['new']}:${fcol(NF)}${ch_rows['new'] + nch - 1},{chm},{f})"
                       f"*IF($C{rr}=\"prepaid_3m\",INDEX(${fcol(1)}${ch_rows['mixp']}:${fcol(NF)}${ch_rows['mixp'] + nch - 1},"
                       f"{chm},{f}),1-INDEX(${fcol(1)}${ch_rows['mixp']}:${fcol(NF)}${ch_rows['mixp'] + nch - 1},{chm},{f}))")
                ws[h] = (f"=IF({k}<1,0,MIN(1,MAX(0,(1-INDEX({s_tbl},$E{rr},{k}+1))*IF({k}<={tw},1-{c}${dr['titr']},1)"
                         f"*{c}${dr['chp']}-{c}${dr['dun']}*{flag})))")
                ws[a] = f"=IF({k}<0,0,IF({k}=0,{new},{prev_a}*(1-{h}*(1-{c}${dr['pause']}))))"
                ws[p] = f"=IF({k}<1,0,{prev_a}*{h}*{c}${dr['pause']})"
                ws[s] = f"=({a}-{p})*$G{rr}"
                ws[f"{bc('C', f)}{rr}"] = f"={s}*{flag}"
                ws[f"{bc('D', f)}{rr}"] = f"={s}*INDEX({R['dose']},1,MIN(MAX({k}+1,1),MAX({R['dose_tiers']})))"

        # summary formulas
        m.header(ws, sum_start - 1, ["Forecast summary", "", ""] + [f"=Levers!{fcol(f)}${self.fm_row}" for f in range(1, NF + 1)])
        col_rng = lambda b, f: f"{bc(b, f)}${seg0}:{bc(b, f)}${seg1}"
        A_, C_, P_ = f"$A${seg0}:$A${seg1}", f"$C${seg0}:$C${seg1}", None
        lam = R["fc_lambda"]
        labels_sum = {
            "new_total": "New paying customers", "active": "Active customers (incl. paused)", "paused": "Paused (lever)",
            "shipments": "Paid shipments", "ship_k0": "First-month shipments", "charges_m": "Charges: monthly plan",
            "charges_p": "Charges: prepaid", "first_m": "First charges: monthly", "first_p": "First charges: prepaid",
            "charges_p_new": "Prepaid charges: new forecast cohorts", "ship_p_new": "Prepaid shipments: new forecast cohorts",
            "disc_promo": "Discount: first-month promo", "disc_prepaid": "Discount: extra prepaid offer (lever c)",
            "promo_first": "Promo first charges", "ship_prepaid": "Prepaid paid shipments",
            "dose_ship": "Dose-weighted shipments", "billing_gross": "Gross billings", "discount": "Discounts",
            "billing_net": "Net billings", "first_net": "Net billings: first orders", "renew_net": "Net billings: renewals",
            "refunds": "Refunds", "chargebacks": "Chargebacks", "recognized": "Recognized revenue (prepaid ratable)",
            "net_revenue": "Net revenue", "med": "Medication (COGS)", "disp": "Pharmacy dispensing",
            "consult": "Clinical consults", "shipping": "Shipping", "packaging": "Packaging / fulfillment",
            "processing": "Payment processing", "cbfees": "Chargeback fees", "tvc": "Total variable cost",
            "contribution": "Contribution profit", "margin": "Contribution margin %", "marketing": "Marketing spend",
            "opex": "Fixed opex", "op_profit": "Operating profit", "card": "Card payouts", "cash_refunds":
            "Refunds + chargebacks paid", "receipts": "Cash receipts", "pharm_paid": "Pharmacy paid (net-30)",
            "mkt_paid": "Marketing paid (lagged)", "other_paid": "Other variable paid", "opex_paid": "Fixed opex paid",
            "payments": "Cash payments", "ncf": "Net operating cash flow", "opening": "Opening cash",
            "closing": "Closing cash"}
        mk_act = f"Monthly_PnL!$B${m.P['mkt']}:${L(1 + self.nm)}${m.P['mkt']}"
        close_act = f"Cash_Flow!$B${m.C['close']}:${L(1 + self.nm)}${m.C['close']}"
        for nm_, rr in S.items():
            bold = nm_ in ("net_revenue", "contribution", "op_profit", "receipts", "payments", "ncf", "closing")
            ws.cell(rr, 1, labels_sum[nm_]).font = m.BOLD if bold else Font()
            for f in range(1, NF + 1):
                c = fcol(f)
                pc = fcol(f - 1) if f > 1 else None
                v = lambda key: f"{c}{S[key]}"
                fx = {
                    "new_total": f"SUM({c}{ch_rows['new']}:{c}{ch_rows['new'] + nch - 1})",
                    "active": f"SUM({col_rng('A', f)})",
                    "paused": f"SUM({col_rng('P', f)})",
                    "shipments": f"SUM({col_rng('S', f)})",
                    "ship_k0": f"SUMIFS({col_rng('S', f)},{A_},{c}$21)",
                    "charges_m": f"SUMIFS({col_rng('C', f)},{C_},\"monthly\")",
                    "charges_p": f"SUMIFS({col_rng('C', f)},{C_},\"prepaid_3m\")",
                    "first_m": f"SUMIFS({col_rng('C', f)},{C_},\"monthly\",{A_},{c}$21)",
                    "first_p": f"SUMIFS({col_rng('C', f)},{C_},\"prepaid_3m\",{A_},{c}$21)",
                    "charges_p_new": f"SUMIFS({col_rng('C', f)},{C_},\"prepaid_3m\",{A_},\">\"&{n})",
                    "ship_p_new": f"SUMIFS({col_rng('S', f)},{C_},\"prepaid_3m\",{A_},\">\"&{n})",
                    "promo_first": f"SUM({c}{ch_rows['promo1']}:{c}{ch_rows['promo1'] + nch - 1})",
                    "ship_prepaid": f"SUMIFS({col_rng('S', f)},{C_},\"prepaid_3m\")",
                    "dose_ship": f"SUM({col_rng('D', f)})",
                    "billing_gross": f"{v('charges_m')}*{c}${dr['price_m']}+{v('charges_p')}*{c}${dr['price_p']}",
                    "disc_promo": f"{v('promo_first')}*{c}${dr['price_m']}*{R['promo_disc']}",
                    "disc_prepaid": f"{v('charges_p_new')}*{c}${dr['price_p']}*{c}${dr['pdisc']}",
                    "discount": f"{v('disc_promo')}+{v('disc_prepaid')}",
                    "billing_net": f"{v('billing_gross')}-{v('discount')}",
                    "first_net": (f"{v('first_m')}*{c}${dr['price_m']}-{v('disc_promo')}"
                                  f"+{v('first_p')}*{c}${dr['price_p']}*(1-{c}${dr['pdisc']})"),
                    "renew_net": f"{v('billing_net')}-{v('first_net')}",
                    "refunds": f"({v('first_net')}*{R['ref_first']}+{v('renew_net')}*{R['ref_renew']})*{R['ref_share']}",
                    "chargebacks": f"{v('first_net')}*{R['cb_first']}+{v('renew_net')}*{R['cb_renew']}",
                    "recognized": (f"{v('charges_m')}*{c}${dr['price_m']}-{v('disc_promo')}"
                                   f"+({v('ship_prepaid')}-{v('ship_p_new')}*{c}${dr['pdisc']})*{c}${dr['price_p']}/{R['term_p']}"),
                    "net_revenue": f"{v('recognized')}-{v('refunds')}-{v('chargebacks')}",
                    "med": f"{v('dose_ship')}*{c}${dr['med']}",
                    "disp": f"{v('shipments')}*{R['disp']}",
                    "consult": f"{v('ship_k0')}*{R['consult_alloc']}+({v('shipments')}-{v('ship_k0')})*{R['consult_fu']}",
                    "shipping": f"{v('shipments')}*{c}${dr['ship']}",
                    "packaging": f"{v('shipments')}*{c}${dr['pack']}",
                    "processing": f"{v('billing_net')}*{R['proc_pct']}+({v('charges_m')}+{v('charges_p')})*{R['proc_fixed']}",
                    "cbfees": (f"(({v('first_m')}+{v('first_p')})*{R['cb_first']}+({v('charges_m')}+{v('charges_p')}"
                               f"-{v('first_m')}-{v('first_p')})*{R['cb_renew']})*{R['cb_fee']}"),
                    "tvc": f"SUM({c}{S['med']}:{c}{S['cbfees']})",
                    "contribution": f"{v('net_revenue')}-{v('tvc')}",
                    "margin": f"IFERROR({v('contribution')}/{v('net_revenue')},0)",
                    "marketing": f"SUM({c}{ch_rows['spend']}:{c}{ch_rows['spend'] + nch - 1})",
                    "opex": f"Levers!{c}${self.opex_row}",
                    "op_profit": f"{v('contribution')}-{v('marketing')}-{v('opex')}",
                    "card": (f"{v('billing_net')}*(1-{lam})+" +
                             (R["fc_carry_card"] if f == 1 else f"{pc}{S['billing_net']}*{lam}")),
                    "cash_refunds": f"{v('refunds')}+{v('chargebacks')}",
                    "receipts": f"{v('card')}-{v('cash_refunds')}",
                    "pharm_paid": R["fc_carry_pharm"] if f == 1 else f"{pc}{S['med']}+{pc}{S['disp']}",
                    "mkt_paid": (f"IF({c}$20-{R['ad_lag']}>=1,INDEX($D${S['marketing']}:${fcol(NF)}${S['marketing']},1,"
                                 f"{c}$20-{R['ad_lag']}),INDEX({mk_act},1,{n}+{c}$20-{R['ad_lag']}))"),
                    "other_paid": f"{v('consult')}+{v('shipping')}+{v('packaging')}+{v('processing')}+{v('cbfees')}",
                    "opex_paid": v("opex"),
                    "payments": f"{v('pharm_paid')}+{v('mkt_paid')}+{v('other_paid')}+{v('opex_paid')}",
                    "ncf": f"{v('receipts')}-{v('payments')}",
                    "opening": f"INDEX({close_act},1,{n})" if f == 1 else f"{pc}{S['closing']}",
                    "closing": f"{v('opening')}+{v('ncf')}",
                }[nm_]
                cell = ws.cell(rr, 3 + f, "=" + fx)
                cell.number_format = ("0.0%" if nm_ == "margin" else "#,##0" if rr < S["billing_gross"] else m.USD)
                if bold:
                    cell.font = m.BOLD
        ws.column_dimensions["A"].width = 44
        for j in range(2, 3 + NF + 1):
            ws.column_dimensions[L(j)].width = 14
        ws.freeze_panes = "D20"
        return S

    # ------------------------------------------------------------------ build all
    def build(self):
        self.levers()
        self.scenario_params()
        master = [f"=1" for _ in LEVER_KEYS]          # Levers!B (On) is applied inside the on-factor
        self.S = self.engine("Forecast", f"={self.R['scenario']}", master,
                             "Forecast: 9-month driver-based cohort roll-forward (scenario per Inputs switch, all levers per Levers)")
        self.engine("Eng_Base", "=1", master, "Engine: base case, all levers")
        self.engine("Eng_Down", "=2", master, "Engine: downside case, all levers")
        for i in range(len(LEVER_KEYS)):
            flags = ["=1" if j < i else "=0" for j in range(len(LEVER_KEYS))]
            self.engine(f"Eng_Bridge{i}", f"={self.R['scenario']}", flags,
                        f"Engine: bridge step {i} (levers 1..{i} on, selected scenario)")
        self.scenarios_tables()
        self.bridge()

    # ------------------------------------------------------------------ Scenarios tables
    def scenarios_tables(self):
        ws, m, S, R = self.wb["Scenarios"], self.m, self.S, self.R
        metrics = [("New customers", "new_total", "#,##0"), ("Active customers", "active", "#,##0"),
                   ("Net revenue", "net_revenue", m.USD), ("Contribution", "contribution", m.USD),
                   ("Operating profit", "op_profit", m.USD), ("Net operating cash flow", "ncf", m.USD),
                   ("Closing cash", "closing", m.USD)]
        r = 11
        ws.cell(r, 1, "Baseline month").font = m.BOLD
        ws.cell(r, 2, f"={R['baseline_month']}")
        r += 1
        ws.cell(r, 1, "Baseline net operating cash flow").font = m.BOLD
        ws.cell(r, 2, (f"=INDEX(Cash_Flow!$B${m.C['ncf']}:${L(1 + self.nm)}${m.C['ncf']},1,"
                       f"MATCH({R['baseline_month']},{R['months']},0))")).number_format = m.USD
        self.base_ncf = f"Scenarios!$B${r}"
        r += 1
        ws.cell(r, 1, "Target (multiple x baseline)").font = m.BOLD
        ws.cell(r, 2, f"={R['target_mult']}*{self.base_ncf}").number_format = m.USD
        self.target = f"Scenarios!$B${r}"
        r += 2
        self.sc = {}
        for j, (eng, lab) in enumerate([("Eng_Base", "Base case"), ("Eng_Down", "Downside case")]):
            c0 = 2 + j * (NF + 1)
            ws.cell(r, c0, lab).font = m.BOLD
            for f in range(1, NF + 1):
                ws.cell(r + 1, c0 + f - 1, f"=Levers!{fcol(f)}${self.fm_row}").font = m.BOLD
            rr = r + 2
            rows = {}
            for lab_, key, fmt in metrics:
                ws.cell(rr, 1, lab_)
                for f in range(1, NF + 1):
                    ws.cell(rr, c0 + f - 1, f"={eng}!{fcol(f)}{S[key]}").number_format = fmt
                rows[key] = rr
                rr += 1
            # sustained flag row
            ws.cell(rr, 1, "2x reached and sustained from this month")
            last = L(c0 + NF - 1)
            for f in range(1, NF + 1):
                col = L(c0 + f - 1)
                ws.cell(rr, c0 + f - 1, f"=MIN({col}{rows['ncf']}:{last}{rows['ncf']})>={self.target}")
            rows["flag"] = rr
            self.sc[eng] = (c0, rows)
        r = rr + 2
        m.header(ws, r, ["Summary", "Base case", "Downside case"])
        r += 1
        labels = ["Month 2x first reached and sustained", "Peak additional cash needed (opening - lowest closing)",
                  "Month of lowest closing cash", "Net operating cash flow, last forecast month",
                  "Gap to target, last forecast month", "Funding need (cash below zero)"]
        self.sum_rows = {}
        for i, lab in enumerate(labels):
            ws.cell(r + i, 1, lab)
            self.sum_rows[lab] = r + i
        for j, eng in enumerate(["Eng_Base", "Eng_Down"]):
            c0, rows = self.sc[eng]
            first, last = L(c0), L(c0 + NF - 1)
            months = f"${first}${rows['flag'] - len(metrics) - 1}:${last}${rows['flag'] - len(metrics) - 1}"
            close = f"${first}${rows['closing']}:${last}${rows['closing']}"
            ncf = f"${first}${rows['ncf']}:${last}${rows['ncf']}"
            col = 2 + j
            vals = [
                f"=IFERROR(INDEX({months},1,MATCH(TRUE,${first}${rows['flag']}:${last}${rows['flag']},0)),\"not reached\")",
                f"=MAX(0,{eng}!$D${S['opening']}-MIN({close}))",
                f"=INDEX({months},1,MATCH(MIN({close}),{close},0))",
                f"=INDEX({ncf},1,{NF})",
                f"={self.target}-INDEX({ncf},1,{NF})",
                f"=MAX(0,-MIN({close}))",
            ]
            for i, v in enumerate(vals):
                c = ws.cell(r + i, col, v)
                c.number_format = m.USD
        ws.column_dimensions["A"].width = 52
        for jj in range(2, 2 + 2 * (NF + 1)):
            ws.column_dimensions[L(jj)].width = 14

    # ------------------------------------------------------------------ Bridge
    def bridge(self):
        ws, m, S, R = self.wb["Doubling_Bridge"], self.m, self.S, self.R
        ws["A1"] = "Doubling bridge: baseline month cash flow to target-month cash flow, one bar per lever"
        ws["A1"].font = m.TITLE
        ws["A2"] = ("Order (fixed): momentum with no levers, then levers switched on cumulatively a, b1, b2, b3, c, d, e, f. "
                    "Each bar = change in the target month's net operating cash flow when that lever is added. "
                    "Bars sum exactly to the total change.")
        ws["A2"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[2].height = 30
        ws["A4"], ws["B4"] = "Bridge target month", f"={R['bridge_month']}"
        ws["A5"], ws["B5"] = "Target month # (forecast)", f"=MATCH(B4,{self.fm_rng},0)"
        ws["A6"], ws["B6"] = "Scenario # (from Inputs switch)", f"={R['scenario']}"
        ncf_at = lambda eng: f"INDEX({eng}!$D${S['ncf']}:${fcol(NF)}${S['ncf']},1,$B$5)"
        con_at = lambda eng: f"INDEX({eng}!$D${S['contribution']}:${fcol(NF)}${S['contribution']},1,$B$5)"
        ws["A7"], ws["B7"] = "Baseline net operating cash flow", f"={self.base_ncf}"
        ws["A8"], ws["B8"] = "Target (2x baseline)", f"={self.target}"
        ws["A9"], ws["B9"] = "Target-month cash flow (selected scenario, all levers)", f"={ncf_at('Forecast')}"
        ws["A10"], ws["B10"] = "Total change", "=B9-B7"
        for rr in range(7, 11):
            ws[f"B{rr}"].number_format = m.USD
        m.header(ws, 12, ["Step", "Cash impact $/month", "Contribution impact $/month", "% of uplift", "Timing",
                          "Confidence", "Evidence needed", "Concentration risk", "Waterfall base", "Increase",
                          "Decrease"])
        r = 13
        ws.cell(r, 1, "Baseline")
        ws.cell(r, 9, 0)
        ws.cell(r, 10, "=B7")
        ws.cell(r, 11, 0)
        r += 1
        steps = [("Momentum (no levers)", "Eng_Bridge0", None, None)]
        engines = [f"Eng_Bridge{i}" for i in range(1, len(LEVER_KEYS))] + ["Forecast"]
        prev = "Eng_Bridge0"
        for key, eng in zip(LEVER_KEYS, engines):
            steps.append((None, eng, prev, key))
            prev = eng
        self.bar_rows = []
        first = r
        for label, eng, prev_eng, key in steps:
            if key is None:
                ws.cell(r, 1, label)
                ws.cell(r, 2, f"={ncf_at(eng)}-B7")
                ws.cell(r, 3, f"={con_at(eng)}-INDEX(Monthly_PnL!$B${m.P['contrib']}:${L(1 + self.nm)}${m.P['contrib']},1,"
                              f"MATCH({R['baseline_month']},{R['months']},0))")
                ws.cell(r, 5, "organic growth of the active base at constant spend")
                ws.cell(r, 6, "medium")
                ws.cell(r, 7, "Retention curves hold for new cohorts; CPA stable at current spend.")
            else:
                lr = self.lv_rows[key]
                ws.cell(r, 1, f"=Levers!A{lr}")
                ws.cell(r, 2, f"={ncf_at(eng)}-{ncf_at(prev_eng)}")
                ws.cell(r, 3, f"={con_at(eng)}-{con_at(prev_eng)}")
                ws.cell(r, 5, f"=\"from \"&Levers!E{lr}&\", \"&Levers!F{lr}&\"-month ramp\"")
                ws.cell(r, 6, f"=Levers!G{lr}")
                ws.cell(r, 7, f"=Levers!I{lr}")
                ws.cell(r, 8, f"=IF(D{r}>0.4,\"CONCENTRATION RISK (>40% of uplift)\",\"\")")
            ws.cell(r, 4, f"=IFERROR(B{r}/$B$10,0)")
            ws.cell(r, 9).value = f"=IF(B{r}>=0,J$13+SUM(B${first}:B{r})-B{r},J$13+SUM(B${first}:B{r}))"
            ws.cell(r, 10, f"=MAX(B{r},0)")
            ws.cell(r, 11, f"=MAX(-B{r},0)")
            self.bar_rows.append(r)
            r += 1
        last = r - 1
        ws.cell(r, 1, "Target-month cash flow").font = m.BOLD
        ws.cell(r, 2, f"=B7+SUM(B{first}:B{last})").font = m.BOLD
        ws.cell(r, 9, 0)
        ws.cell(r, 10, f"=B{r}")
        ws.cell(r, 11, 0)
        self.bridge_total_row = r
        r += 1
        ws.cell(r, 1, "Check: bars sum to total change")
        ws.cell(r, 2, f"=SUM(B{first}:B{last})-B10")
        ws.cell(r, 3, f"=MAX(B{r},-B{r})<0.000001")
        for rr in range(13, r + 1):
            for col in (2, 3, 9, 10, 11):
                ws.cell(rr, col).number_format = m.USD
            ws.cell(rr, 4).number_format = "0.0%"
            ws.cell(rr, 7).alignment = Alignment(wrap_text=True, vertical="top")
        self.bridge_check = f"Doubling_Bridge!$C${r}"
        chart = BarChart()
        chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
        chart.title = "Baseline to target-month net operating cash flow"
        data = Reference(ws, min_col=9, max_col=11, min_row=12, max_row=self.bridge_total_row)
        cats = Reference(ws, min_col=1, min_row=13, max_row=self.bridge_total_row)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.series[0].graphicalProperties.noFill = True
        chart.series[0].graphicalProperties.line.noFill = True
        chart.series[1].graphicalProperties.solidFill = "4E9A6B"
        chart.series[2].graphicalProperties.solidFill = "C0504D"
        chart.height, chart.width = 9, 22
        ws.add_chart(chart, f"A{r + 3}")
        ws.column_dimensions["A"].width = 50
        for col, w in zip("BCDEFGHIJK", [16, 16, 10, 28, 11, 60, 30, 14, 14, 14]):
            ws.column_dimensions[col].width = w
        ws.freeze_panes = "B13"
