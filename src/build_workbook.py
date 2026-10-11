"""Build output/omzo_model.xlsx: a live-formula financial model (Excel + Google Sheets).

Design:
- Data_* sheets hold VOLUMES aggregated from data/*.csv (orders, shipments, payments by
  day, refunds, customers). Every price, cost rate and timing rule is applied by a live
  formula pointing at Inputs. Editing an input therefore re-prices the P&L, cash flow
  and unit economics. Behavioral inputs (churn, pause, refund rates...) shaped the
  simulated volumes; changing them does not re-simulate data (flagged on Inputs).
- Formulas use only SUM, SUMIFS, COUNTIFS, INDEX, MATCH, IF, IFERROR, ROUND, MIN, MAX
  plus arithmetic, with absolute references. No named ranges, tables or array spills.

Usage:  python src/build_workbook.py [--out output/omzo_model.xlsx]
"""
from __future__ import annotations

import argparse
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import ltv_projection
from forecast_sheets import NF, Forecast
from unit_economics import ROOT, load_assumptions

DATA = ROOT / "data"
METRICS = ROOT / "output" / "metrics"
OUT = ROOT / "output" / "omzo_model.xlsx"

BLUE = Font(color="0000FF")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
HEAD_FILL = PatternFill("solid", fgColor="DDE3EA")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
USD = '"$"#,##0;[Red]-"$"#,##0'
USD2 = '"$"#,##0.00;[Red]-"$"#,##0.00'
PCT = "0.0%"
INT = "#,##0"
NUM2 = "0.00"
DATE_FMT = "yyyy-mm-dd"
N_MONTH_COLS = 9          # columns B..J on monthly sheets (one per modeled month)


def mcol(m: int) -> str:
    """Column letter of month m (1-based) on monthly sheets: B..J."""
    return L(1 + m)


def absr(sheet: str, col: str, row: int) -> str:
    return f"{sheet}!${col}${row}"


def rng(sheet: str, col: str, r1: int, r2: int) -> str:
    return f"{sheet}!${col}${r1}:${col}${r2}"


def header(ws, row, values, start_col=1):
    for j, v in enumerate(values):
        c = ws.cell(row, start_col + j, v)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(horizontal="center", wrap_text=True)


def section(ws, row, text, width=11):
    ws.cell(row, 1, text).font = BOLD
    for j in range(1, width + 1):
        ws.cell(row, j).fill = SECTION_FILL


def widths(ws, first=34, rest=14, n=20):
    ws.column_dimensions["A"].width = first
    for j in range(2, n + 1):
        ws.column_dimensions[L(j)].width = rest


# ============================================================================= model
class Model:
    def __init__(self, a: dict, scenario: int = 1):
        self.a, self.scenario = a, scenario
        self.TITLE, self.BLUE, self.BOLD, self.USD, self.USD2 = TITLE, BLUE, BOLD, USD, USD2
        self.header, self.section = header, section
        self.months = [str(m) for m in a["period"]["months"]]
        self.nm = len(self.months)
        self.channels = list(a["channels"])
        self.wb = Workbook()
        self.R: dict[str, str] = {}          # input name -> absolute cell
        self.D: dict[str, dict[str, str]] = {}  # data sheet -> column name -> range
        self.load_data()

    # ------------------------------------------------------------------ data prep
    def load_data(self):
        rd = lambda n, dates=(): pd.read_csv(DATA / f"{n}.csv", parse_dates=list(dates))
        self.cust = rd("customers", ["signup_date"])
        self.orders = rd("orders", ["order_date"])
        self.pay = rd("payments", ["attempt_date", "payout_date"])
        self.ref = rd("refunds", ["refund_date"])
        self.ship = rd("shipments", ["ship_date"])
        self.oc = rd("order_costs")
        self.midx = {m: i + 1 for i, m in enumerate(self.months)}
        self.idx_of = lambda s: s.dt.strftime("%Y-%m").map(self.midx)

    # ------------------------------------------------------------------ build
    def build(self) -> Workbook:
        wb = self.wb
        wb.active.title = "README"
        for name in ["Inputs", "Monthly_PnL", "Cash_Flow", "Levers", "Scenarios", "Doubling_Bridge",
                     "Unit_Economics", "Channels", "Cohorts", "Calc_LTV", "Checks"]:
            wb.create_sheet(name)
        self.inputs()
        self.data_sheets()
        self.pnl()
        self.cash_flow()
        self.channels_sheet()
        self.calc_ltv()
        self.unit_economics()
        self.cohorts()
        self.fc = Forecast(self)
        self.fc.build()
        self.checks()
        self.readme()
        order = ["README", "Inputs", "Monthly_PnL", "Cash_Flow", "Forecast", "Levers", "Scenarios",
                 "Doubling_Bridge", "Unit_Economics", "Channels", "Cohorts", "Calc_LTV", "Checks"]
        rest = [ws for ws in wb.worksheets if ws.title not in order]
        wb._sheets = [wb[n] for n in order] + sorted(rest, key=lambda w: (w.title.startswith("Eng_"), w.title))
        return wb

    # ------------------------------------------------------------------ Inputs
    def inputs(self):
        a, ws, R = self.a, self.wb["Inputs"], self.R
        ws["A1"] = "Inputs: every assumption from assumptions.yaml (blue = editable input)"
        ws["A1"].font = TITLE
        ws["A2"] = ("'Drives' column: model = re-prices formulas live; data = shaped the simulated "
                    "volumes (edit + re-run src/generate_data.py to change the data); info = reference only.")
        ws["A2"].font = Font(italic=True)
        r = 4

        def scalar(key, label, value, unit, drives="model", fmt=None):
            nonlocal r
            ws.cell(r, 1, label)
            c = ws.cell(r, 2, value)
            c.font = BLUE
            if fmt:
                c.number_format = fmt
            ws.cell(r, 3, unit)
            ws.cell(r, 4, drives)
            R[key] = absr("Inputs", "B", r)
            r += 1

        def series(key, label, values, unit, drives="model", fmt=None, start_col=2):
            nonlocal r
            ws.cell(r, 1, label)
            for j, v in enumerate(values):
                c = ws.cell(r, start_col + j, v)
                c.font = BLUE
                if fmt:
                    c.number_format = fmt
            ws.cell(r, start_col + len(values), unit)
            ws.cell(r, start_col + len(values) + 1, drives)
            R[key] = f"Inputs!${L(start_col)}${r}:${L(start_col + len(values) - 1)}${r}"
            R[key + "_row"] = r
            r += 1

        section(ws, r, "Period"); r += 1
        header(ws, r, ["Month #"] + list(range(1, self.nm + 1)))
        R["month_idx_row"] = f"Inputs!$B${r}:${L(1 + self.nm)}${r}"
        r += 1
        series("months", "Month", self.months, "YYYY-MM", "info")
        R["n_months"] = f"MAX({R['month_idx_row']})"
        scalar("seed", "Random seed (data generator)", a["period"]["random_seed"], "", "data")
        r += 1

        section(ws, r, "Plans & pricing"); r += 1
        p = a["plans"]
        scalar("price_m", "Monthly plan price", p["monthly"]["price"], "$ per charge", fmt=USD2)
        scalar("term_m", "Monthly plan term", p["monthly"]["term_months"], "months per charge")
        scalar("mix_m", "Monthly plan mix (new customers)", p["monthly"]["mix"], "share", "data", PCT)
        scalar("price_p", "3-month prepaid price", p["prepaid_3m"]["price"], "$ per charge", fmt=USD2)
        scalar("term_p", "3-month prepaid term", p["prepaid_3m"]["term_months"], "months per charge")
        scalar("mix_p", "Prepaid plan mix (new customers)", p["prepaid_3m"]["mix"], "share", "data", PCT)
        scalar("prepaid_churn_mult", "Prepaid churn multiplier", p["prepaid_3m"]["churn_multiplier"], "x", "data")
        scalar("promo_disc", "First-charge promo discount (monthly plan)", a["promo"]["first_month_discount"],
               "share of price", fmt=PCT)
        r += 1

        section(ws, r, "Channel settings"); r += 1
        header(ws, r, ["Channel", "Paid media (1/0)", "Promo (1/0)", "Expected CPA ($)",
                       "Churn multiplier", "Period spend ($)"])
        r += 1
        ch0 = r
        for ch, cfg in a["channels"].items():
            ws.cell(r, 1, ch)
            for j, v in enumerate([int(cfg["paid"]), int(ch in a["promo"]["channels"]),
                                   cfg["expected_cpa"], cfg["churn_multiplier"]]):
                c = ws.cell(r, 2 + j, v)
                c.font = BLUE
            ws.cell(r, 4).number_format = USD2
            r += 1
        ch1 = r - 1
        R["ch_names"] = rng("Inputs", "A", ch0, ch1)
        R["ch_paid"] = rng("Inputs", "B", ch0, ch1)
        R["ch_promo"] = rng("Inputs", "C", ch0, ch1)
        R["ch_cpa_exp"] = rng("Inputs", "D", ch0, ch1)
        R["ch_spend_total"] = rng("Inputs", "F", ch0, ch1)
        settings_rows = dict(zip(self.channels, range(ch0, ch1 + 1)))
        r += 1

        section(ws, r, "Marketing spend by channel and month ($, paid media + referral credits)"); r += 1
        header(ws, r, ["Channel"] + self.months + ["Total"])
        r += 1
        sp0 = r
        for ch, cfg in a["channels"].items():
            ws.cell(r, 1, ch)
            for j, v in enumerate(cfg["spend"]):
                c = ws.cell(r, 2 + j, v)
                c.font, c.number_format = BLUE, USD
            t = ws.cell(r, 2 + self.nm, f"=SUM(B{r}:{mcol(self.nm)}{r})")
            t.number_format = USD
            ws.cell(settings_rows[ch], 6, f"={L(2 + self.nm)}{r}").number_format = USD
            r += 1
        R["spend_r0"], R["spend_r1"] = sp0, r - 1
        r += 1

        section(ws, r, "Retention dynamics (drive the simulated volumes)"); r += 1
        rt = a["retention"]
        header(ws, r, ["Tenure month"] + [str(t) for t in range(1, len(rt["churn_by_tenure"]))]
               + [f"{len(rt['churn_by_tenure'])}+"])
        r += 1
        series("churn_tenure", "Voluntary churn by tenure month", rt["churn_by_tenure"], "share", "data", PCT)
        scalar("pause_rate", "Monthly pause rate", rt["monthly_pause_rate"], "share of actives", "data", PCT)
        scalar("pause_dur", "Pause duration", rt["pause_duration_months"], "months", "data")
        scalar("react_rate", "Win-back (reactivation) rate", rt["reactivation_rate"], "monthly share of churned", "data", PCT)
        scalar("react_window", "Win-back window", rt["reactivation_window_months"], "months", "data")
        scalar("fail_rate", "Failed renewal payment rate", rt["failed_payment_rate"], "share of renewals", "data", PCT)
        scalar("fail_recovery", "Failed payment recovery (dunning)", rt["failed_payment_recovery"], "share", "data", PCT)
        r += 1

        section(ws, r, "Medication (COGS): the only medication cost input"); r += 1
        header(ws, r, ["Month"] + self.months)
        r += 1
        series("med_base", "Medication cost per 1-month supply", a["medication"]["cost_per_month_supply"],
               "$ at starting dose", fmt=USD2)
        dose = a["medication"]["dose_multiplier_by_tenure"]
        header(ws, r, ["Dose tier (tenure month)"] + list(range(1, len(dose) + 1)))
        R["dose_tiers"] = f"Inputs!$B${r}:${L(1 + len(dose))}${r}"
        r += 1
        series("dose", "Dose multiplier by tenure month (last tier = and later)", dose, "x")
        r += 1

        section(ws, r, "Variable costs (per unit; exclude medication)"); r += 1
        vc = a["variable_costs"]
        scalar("disp", "Pharmacy dispensing", vc["pharmacy_dispensing_per_shipment"], "$ per shipment", fmt=USD2)
        scalar("consult_init", "Clinical consult: initial", vc["clinical_consult"]["initial"], "$ per applicant", fmt=USD2)
        scalar("consult_fu", "Clinical consult: follow-up", vc["clinical_consult"]["follow_up"], "$ per shipment (tenure 2+)", fmt=USD2)
        scalar("decline", "Clinical intake decline rate", vc["clinical_consult"]["intake_decline_rate"],
               "share (allocated to approved)", fmt=PCT)
        ws.cell(r, 1, "Initial consult per approved customer (calc)")
        ws.cell(r, 2, f"=ROUND({R['consult_init']}/(1-{R['decline']}),2)").number_format = USD2
        ws.cell(r, 3, "$ (rounded to cents)")
        ws.cell(r, 4, "formula")
        R["consult_alloc"] = absr("Inputs", "B", r)
        r += 1
        scalar("shipping", "Shipping", vc["shipping_per_shipment"], "$ per shipment", fmt=USD2)
        scalar("packaging", "Packaging / fulfillment", vc["packaging_fulfillment_per_shipment"], "$ per shipment", fmt=USD2)
        scalar("proc_pct", "Payment processing %", vc["payment_processing"]["pct_of_charge"], "of charge", fmt="0.00%")
        scalar("proc_fixed", "Payment processing fixed fee", vc["payment_processing"]["fixed_per_charge"], "$ per charge", fmt=USD2)
        r += 1

        section(ws, r, "Refunds & chargebacks"); r += 1
        rf, cb = a["refunds"], a["chargebacks"]
        scalar("ref_first", "Refund rate: first orders", rf["first_order_rate"], "share of charges", "data*", PCT)
        scalar("ref_renew", "Refund rate: renewals", rf["renewal_rate"], "share of charges", "data*", PCT)
        scalar("ref_share", "Average refund share", rf["avg_refund_share"], "share of charge", "data*", PCT)
        scalar("ref_lag", "Refund lag", rf["lag_months"], "months", "data")
        scalar("cb_first", "Chargeback rate: first orders", cb["first_order_rate"], "share of charges", "data*", PCT)
        scalar("cb_renew", "Chargeback rate: renewals", cb["renewal_rate"], "share of charges", "data*", PCT)
        scalar("cb_fee", "Chargeback fee", cb["fee_per_chargeback"], "$ per chargeback", fmt=USD2)
        scalar("cb_lag", "Chargeback lag", cb["lag_months"], "months", "data")
        ws.cell(r, 1, "data* = shaped the data; also used live by the LTV projection").font = Font(italic=True)
        r += 2

        section(ws, r, "Fixed operating expenses ($ per month)"); r += 1
        header(ws, r, ["Category"] + self.months)
        r += 1
        ox0 = r
        for cat, vals in a["fixed_opex"].items():
            series(f"opex_{cat}", cat, vals, "$", fmt=USD)
        R["opex_r0"], R["opex_r1"] = ox0, r - 1
        r += 1

        section(ws, r, "Cash timing"); r += 1
        c = a["cash"]
        scalar("opening_cash", "Opening cash balance (2026-01-01)", c["opening_balance"], "$", fmt=USD)
        scalar("days_per_month", "Days per month (lag simplification)", c["days_per_month"], "days", "data")
        scalar("payout_lag", "Card payout lag", c["card_payout_lag_days"], "days (T+n)")
        scalar("pharm_terms", "Pharmacy payment terms", c["pharmacy_payment_terms_days"], "days after shipment")
        scalar("ad_lag", "Ad spend payment lag", c["ad_spend_payment_lag_months"], "months")
        r += 1

        section(ws, r, "Simulation behavior (data generator only)"); r += 1
        s = a["simulation"]
        scalar("noise", "New-customer volume noise (sd)", s["new_customer_noise_sd"], "share", "data", PCT)
        scalar("max_bill_day", "Max billing anchor day", s["max_billing_day"], "day of month", "data")
        series("retry_days", "Dunning retry days", s["dunning_retry_days"], "days after failure", "data")
        scalar("ref_spread", "Refund share spread", s["refund_share_spread"], "+/- share", "data", PCT)
        scalar("ref_jitter", "Refund date jitter", s["refund_jitter_days"], "+/- days", "data")
        scalar("cb_jitter", "Chargeback date jitter", s["chargeback_jitter_days"], "+/- days", "data")
        scalar("ref_cancel", "Refund triggers cancel", int(s["refund_triggers_cancel"]), "1/0", "data")
        scalar("cb_block", "Chargeback blocks customer", int(s["chargeback_blocks_customer"]), "1/0", "data")
        for reason, w in s["cancel_reasons"].items():
            scalar(f"reason_{reason}", f"Cancel reason weight: {reason}", w, "share", "data", PCT)
        r += 1

        section(ws, r, "LTV projection & CPA targets"); r += 1
        lt = a["ltv"]
        scalar("ltv_horizon", "LTV horizon (structural: Calc_LTV has 12 rows)", lt["horizon_months"], "months", "info")
        scalar("disc_rate", "Monthly discount rate", lt["discount_rate_monthly"], "per month", fmt="0.00%")
        scalar("ltv_start", "Projection start month", lt["projection_start_month"], "YYYY-MM")
        scalar("target_ratio", "Target LTV / CPA", lt["target_ltv_to_cpa"], "x")
        scalar("max_payback", "Max CPA payback window", lt["max_payback_months"], "months")
        scalar("min_fit", "Min customers per retention fit point", lt["min_customers_for_fit"], "customers", "data")
        scalar("sens_med", "Sensitivity: medication uplift", lt["sensitivity"]["medication_uplift"], "share", fmt=PCT)
        scalar("sens_ret", "Sensitivity: retention drop (months 1+)", lt["sensitivity"]["retention_drop_pp"],
               "share points", fmt=PCT)
        r += 1

        section(ws, r, "Forecast & doubling plan"); r += 1
        scalar("scenario", "SCENARIO SWITCH (1 = base, 2 = downside): drives the Forecast sheet", self.scenario, "1 or 2")
        scalar("baseline_month", "Baseline month (see README for the choice)", "2026-09", "YYYY-MM")
        scalar("target_mult", "Target multiple of baseline net operating cash flow", 2, "x")
        scalar("bridge_month", "Bridge target month", "2027-06", "YYYY-MM")
        r += 1

        section(ws, r, "Unit-economics script & model settings"); r += 1
        uc = a["unit_check"]
        scalar("uc_month", "Unit check reference month", uc["reference_month"], "YYYY-MM", "info")
        scalar("uc_horizon", "Unit check payback horizon", uc["payback_horizon_months"], "months", "info")
        scalar("tol", "Check tolerance (Checks sheet)", 0.01, "$", fmt=USD2)

        widths(ws, 46, 13, 14)
        ws.column_dimensions["C"].width = 16
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ Data sheets
    def data_sheet(self, name, df: pd.DataFrame, formulas=None, fmts=None):
        """Write values + per-row formula columns. formulas: {col: f(row, letters) -> str}."""
        formulas, fmts = formulas or {}, fmts or {}
        ws = self.wb.create_sheet(name)
        cols = list(df.columns) + list(formulas)
        letters = {c: L(i + 1) for i, c in enumerate(cols)}
        header(ws, 1, cols)
        n = len(df)
        values = df.to_numpy(dtype=object)
        for i in range(n):
            r = i + 2
            row = [v.item() if isinstance(v, np.generic) else v for v in values[i]]
            row = [None if (isinstance(v, float) and np.isnan(v)) else v for v in row]
            row += [f(r, letters) for f in formulas.values()]
            ws.append(row)
        for c, f in fmts.items():
            col = letters[c]
            for r in range(2, n + 2):
                ws[f"{col}{r}"].number_format = f
        ws.freeze_panes = "A2"
        for i, c in enumerate(cols):
            ws.column_dimensions[L(i + 1)].width = max(12, min(22, len(c) + 3))
        self.D[name] = {c: rng(name, letters[c], 2, n + 1) for c in cols}
        return ws

    def data_sheets(self):
        R, a = self.R, self.a
        cal_n = self.nm + NF + 3
        starts = [pd.Period(self.months[0], "M") + i for i in range(cal_n)]
        cal = pd.DataFrame({"month_idx": range(1, cal_n + 1),
                            "month": [str(p) for p in starts],
                            "month_start": [p.start_time.date() for p in starts],
                            "month_end": [p.end_time.date() for p in starts]})
        self.data_sheet("Data_Calendar", cal, fmts={"month_start": DATE_FMT, "month_end": DATE_FMT})
        starts_rng = self.D["Data_Calendar"]["month_start"]
        price = lambda plan: f'IF({plan}="prepaid_3m",{R["price_p"]},{R["price_m"]})'
        term = lambda plan: f'IF({plan}="prepaid_3m",{R["term_p"]},{R["term_m"]})'
        unit = lambda plan, promo: f'({price(plan)}-{promo}*ROUND({price(plan)}*{R["promo_disc"]},2))'

        # customers (row-level)
        c = self.cust.assign(signup_idx=self.idx_of(self.cust.signup_date),
                             promo=self.cust.promo_applied.astype(int))
        self.data_sheet("Data_Customers", c[["customer_id", "signup_idx", "channel", "plan_type", "promo"]])

        # orders aggregated by month x plan x first x promo
        o = self.orders[self.orders.status != "failed"].copy()
        o["order_idx"] = self.idx_of(o.order_date)
        o["is_first"] = (o.order_number == 1).astype(int)
        o["promo"] = (o.discount > 0).astype(int)
        og = (o.groupby(["order_idx", "plan_type", "is_first", "promo"])
              .agg(orders=("order_id", "size"), net_billed_data=("net_price", "sum")).reset_index())
        self.data_sheet("Data_Orders", og, {
            "unit_net_price": lambda r, c: f"={unit(f'{c['plan_type']}{r}', f'{c['promo']}{r}')}",
            "gross": lambda r, c: f"={c['orders']}{r}*{price(f'{c['plan_type']}{r}')}",
            "discount": lambda r, c: f"={c['gross']}{r}-{c['orders']}{r}*{c['unit_net_price']}{r}",
            "net": lambda r, c: f"={c['gross']}{r}-{c['discount']}{r}",
            "processing": lambda r, c: (f"={c['orders']}{r}*ROUND({c['unit_net_price']}{r}*{R['proc_pct']}"
                                        f"+{R['proc_fixed']},2)"),
            "term": lambda r, c: f"={term(f'{c['plan_type']}{r}')}",
            "last_idx": lambda r, c: f"={c['order_idx']}{r}+{c['term']}{r}-1",
            "gross_pm": lambda r, c: f"={c['gross']}{r}/{c['term']}{r}",
            "discount_pm": lambda r, c: f"={c['discount']}{r}/{c['term']}{r}",
            "net_pm": lambda r, c: f"={c['net']}{r}/{c['term']}{r}",
        }, fmts={k: USD2 for k in ["net_billed_data", "unit_net_price", "gross", "discount", "net",
                                   "processing", "gross_pm", "discount_pm", "net_pm"]})

        # shipments aggregated by day x dose tier
        n_tiers = len(a["medication"]["dose_multiplier_by_tenure"])
        s = self.ship.assign(tier=self.ship.tenure_month.clip(upper=n_tiers))
        sg = (s.groupby(["ship_date", "tier"]).size().rename("shipments").reset_index())
        sg["ship_date"] = sg.ship_date.dt.date
        self.data_sheet("Data_Shipments", sg, {
            "ship_idx": lambda r, c: f"=MATCH({c['ship_date']}{r},{starts_rng},1)",
            "paid_idx": lambda r, c: f"=MATCH({c['ship_date']}{r}+{R['pharm_terms']},{starts_rng},1)",
            "medication": lambda r, c: (f"={c['shipments']}{r}*INDEX({R['med_base']},1,{c['ship_idx']}{r})"
                                        f"*INDEX({R['dose']},1,{c['tier']}{r})"),
            "dispensing": lambda r, c: f"={c['shipments']}{r}*{R['disp']}",
            "consult": lambda r, c: (f"={c['shipments']}{r}*IF({c['tier']}{r}=1,"
                                     f"{R['consult_alloc']},{R['consult_fu']})"),
            "shipping": lambda r, c: f"={c['shipments']}{r}*{R['shipping']}",
            "packaging": lambda r, c: f"={c['shipments']}{r}*{R['packaging']}",
            "pharmacy": lambda r, c: f"={c['medication']}{r}+{c['dispensing']}{r}",
        }, fmts={"ship_date": DATE_FMT, **{k: USD2 for k in ["medication", "dispensing", "consult",
                                                              "shipping", "packaging", "pharmacy"]}})

        # successful payments by attempt day x plan x promo
        p = self.pay[self.pay.status.isin(["success", "recovered"])].merge(
            self.orders[["order_id", "plan_type", "discount"]], on="order_id")
        p["promo"] = (p.discount > 0).astype(int)
        p["attempt_date"] = p.attempt_date.dt.date
        pg = (p.groupby(["attempt_date", "plan_type", "promo"])
              .agg(payments=("payment_id", "size"), amount_data=("amount", "sum")).reset_index())
        self.data_sheet("Data_Payments", pg, {
            "unit_net_price": lambda r, c: f"={unit(f'{c['plan_type']}{r}', f'{c['promo']}{r}')}",
            "amount": lambda r, c: f"={c['payments']}{r}*{c['unit_net_price']}{r}",
            "payout_idx": lambda r, c: f"=MATCH({c['attempt_date']}{r}+{R['payout_lag']},{starts_rng},1)",
        }, fmts={"attempt_date": DATE_FMT, "amount_data": USD2, "unit_net_price": USD2, "amount": USD2})

        # refunds / chargebacks by month x type x first-order flag
        rf = self.ref.merge(self.orders[["order_id", "order_number"]], on="order_id")
        rf["refund_idx"] = self.idx_of(rf.refund_date)
        rf["is_first"] = (rf.order_number == 1).astype(int)
        rg = (rf.groupby(["refund_idx", "type", "is_first"])
              .agg(count=("refund_id", "size"), amount=("amount", "sum")).reset_index())
        self.data_sheet("Data_Refunds", rg, {
            "fee": lambda r, c: f'=IF({c["type"]}{r}="chargeback",{c["count"]}{r}*{R["cb_fee"]},0)',
        }, fmts={"amount": USD2, "fee": USD2})

        # observed LTV inputs by cohort x channel; first orders by channel
        po = self.orders[self.orders.status != "failed"].merge(self.oc[["order_id", "total_variable_cost"]])
        rr = self.ref.merge(self.orders[["order_id", "customer_id", "order_number"]], on="order_id")
        per = self.cust[["customer_id", "channel", "signup_date"]].copy()
        per["cohort_idx"] = self.idx_of(per.signup_date)
        per = per.merge(po.groupby("customer_id").agg(net_charges=("net_price", "sum"),
                                                      order_var_cost=("total_variable_cost", "sum")),
                        on="customer_id", how="left")
        per = per.merge(rr.groupby("customer_id").agg(refunds=("amount", "sum"), cb_fees=("fee", "sum")),
                        on="customer_id", how="left").fillna(0)
        lg = (per.groupby(["cohort_idx", "channel"])
              .agg(customers=("customer_id", "size"), net_charges=("net_charges", "sum"),
                   refunds=("refunds", "sum"), order_var_cost=("order_var_cost", "sum"),
                   cb_fees=("cb_fees", "sum")).reset_index())
        self.data_sheet("Data_LTV", lg, fmts={k: USD2 for k in ["net_charges", "refunds", "order_var_cost", "cb_fees"]})

        f1 = po[po.order_number == 1].merge(self.cust[["customer_id", "channel"]])
        r1 = rr[rr.order_number == 1][["order_id", "amount", "fee"]]
        f1 = f1.merge(r1, on="order_id", how="left").fillna({"amount": 0, "fee": 0})
        fg = (f1.groupby("channel").agg(first_orders=("order_id", "size"), net_charges=("net_price", "sum"),
                                        refunds=("amount", "sum"), order_var_cost=("total_variable_cost", "sum"),
                                        cb_fees=("fee", "sum")).reindex(self.channels).reset_index())
        self.data_sheet("Data_FirstOrders", fg, fmts={k: USD2 for k in ["net_charges", "refunds", "order_var_cost", "cb_fees"]})

        # retention aggregates (activity-based; computed by the SQL layer)
        cl = pd.read_csv(METRICS / "cohort_retention_long.csv").dropna(subset=["retained"])
        cl["cohort_idx"] = cl.cohort.map(self.midx)
        self.data_sheet("Data_Cohorts", cl[["cohort_idx", "months_since_signup", "retained", "retained_paid"]]
                        .astype(int))
        rc = pd.read_csv(METRICS / "renewal_churn.csv")
        rc["month_idx"] = rc.month.map(self.midx)
        self.data_sheet("Data_Renewal", rc[["month_idx", "active_customers", "active_prev_month", "churned",
                                            "renewals_due", "renewed"]].fillna(0).astype(int))
        rbc = pd.read_csv(METRICS / "retention_by_channel.csv")
        self.data_sheet("Data_RetentionChannel", rbc[["channel", "months_since_signup", "eligible_customers",
                                                      "retained", "retained_paid"]])

        # fitted retention curves (Weibull fit in src/ltv_projection.py)
        h = a["ltv"]["horizon_months"]
        rows = []
        for ch in self.channels:
            b, cc = ltv_projection.fit_retention(rbc[rbc.channel == ch], a["ltv"]["min_customers_for_fit"])
            rows.append([ch, b, cc] + list(ltv_projection.retention_curve(b, cc, h)))
        fit = pd.DataFrame(rows, columns=["channel", "fit_b", "fit_c"] + [f"r_k{k}" for k in range(h)])
        self.data_sheet("Data_RetentionFit", fit)

        # forecast inputs: retention by channel x plan (observed where enough customers, Weibull
        # tail beyond, forced non-increasing) and active customers at the end of the actuals
        K = self.nm + NF
        rcp = pd.read_csv(METRICS / "retention_by_channel_plan.csv")
        min_n = a["ltv"]["min_customers_for_fit"]
        rows = []
        for ch in self.channels:
            for plan in ["monthly", "prepaid_3m"]:
                cv = rcp[(rcp.channel == ch) & (rcp.plan_type == plan)].set_index("months_since_signup")
                ok = cv[cv.eligible_customers >= min_n]
                k_last = int(ok.index.max())
                b, cc = ltv_projection.fit_retention(cv.reset_index(), min_n)
                w = ltv_projection.retention_curve(b, cc, K)
                curve = [float(ok.retention_pct.get(k, np.nan)) if k <= k_last
                         else float(ok.retention_pct[k_last] * w[k] / w[k_last]) for k in range(K)]
                curve = list(np.minimum.accumulate(np.nan_to_num(curve, nan=1.0)))
                curve[0] = 1.0
                rows.append([ch, plan, k_last] + curve)
        self.data_sheet("Data_RetentionPlan", pd.DataFrame(
            rows, columns=["channel", "plan_type", "last_observed_k"] + [f"r_k{k}" for k in range(K)]))
        ae = pd.read_csv(METRICS / "active_at_period_end.csv")
        ae["cohort_idx"] = ae.cohort.map(self.midx)
        self.data_sheet("Data_ActiveEnd", ae[["cohort_idx", "channel", "plan_type", "active_customers",
                                              "paying_customers"]])

    # ------------------------------------------------------------------ helpers for monthly sheets
    def month_header(self, ws, title, total=True):
        ws["A1"] = title
        ws["A1"].font = TITLE
        header(ws, 3, ["$ / month"] + self.months + (["Total"] if total else []))
        ws.cell(4, 1, "Month #").font = Font(italic=True, color="808080")
        for m in range(1, self.nm + 1):
            ws.cell(4, 1 + m, f"=INDEX({self.R['month_idx_row']},1,{m})").font = Font(color="808080")
        ws.freeze_panes = "B5"
        widths(ws, 40, 14, 12)

    def month_row(self, ws, r, label, fn, fmt=USD, total="sum", bold=False):
        """fn(col_letter) -> formula for that month column."""
        ws.cell(r, 1, label).font = BOLD if bold else Font()
        for m in range(1, self.nm + 1):
            col = mcol(m)
            c = ws.cell(r, 1 + m, "=" + fn(col))
            c.number_format = fmt
            if bold:
                c.font = BOLD
        tc = ws.cell(r, 2 + self.nm)
        if total == "sum":
            tc.value = f"=SUM(B{r}:{mcol(self.nm)}{r})"
        elif total:
            tc.value = "=" + total
        tc.number_format = fmt
        tc.font = BOLD

    # ------------------------------------------------------------------ P&L
    def pnl(self):
        ws, R, O, F, S = (self.wb["Monthly_PnL"], self.R, self.D["Data_Orders"],
                          self.D["Data_Refunds"], self.D["Data_Shipments"])
        self.month_header(ws, "Monthly P&L (accrual; prepaid recognized ratably; no sales tax)")
        rec = lambda col_rng, col, extra="": (f"SUMIFS({col_rng},{O['order_idx']},\"<=\"&{col}$4,"
                                              f"{O['last_idx']},\">=\"&{col}$4{extra})")
        P = {}
        rows = [
            ("gross", "Gross revenue (recognized)", lambda c: rec(O["gross_pm"], c)),
            ("disc", "Discounts (promo)", lambda c: rec(O["discount_pm"], c)),
            ("refunds", "Refunds", lambda c: f'SUMIFS({F["amount"]},{F["refund_idx"]},{c}$4,{F["type"]},"refund")'),
            ("cbs", "Chargebacks", lambda c: f'SUMIFS({F["amount"]},{F["refund_idx"]},{c}$4,{F["type"]},"chargeback")'),
        ]
        r = 5
        for key, label, fn in rows:
            self.month_row(ws, r, label, fn); P[key] = r; r += 1
        self.month_row(ws, r, "Net revenue", lambda c: f"{c}{P['gross']}-{c}{P['disc']}-{c}{P['refunds']}-{c}{P['cbs']}",
                       bold=True); P["net"] = r; r += 1
        self.month_row(ws, r, "Completed orders", lambda c: f"SUMIFS({O['orders']},{O['order_idx']},{c}$4)", INT)
        P["orders"] = r; r += 2
        section(ws, r, "Variable costs", 2 + self.nm); r += 1
        for key, label, col in [("med", "Medication (COGS)", "medication"),
                                ("disp", "Pharmacy dispensing", "dispensing"),
                                ("consult", "Clinical consults", "consult"),
                                ("ship", "Shipping", "shipping"),
                                ("pack", "Packaging / fulfillment", "packaging")]:
            self.month_row(ws, r, label, lambda c, col=col: f"SUMIFS({S[col]},{S['ship_idx']},{c}$4)")
            P[key] = r; r += 1
        self.month_row(ws, r, "Payment processing", lambda c: f"SUMIFS({O['processing']},{O['order_idx']},{c}$4)")
        P["proc"] = r; r += 1
        self.month_row(ws, r, "Chargeback fees", lambda c: f"SUMIFS({F['fee']},{F['refund_idx']},{c}$4)")
        P["cbfee"] = r; r += 1
        self.month_row(ws, r, "Total variable cost", lambda c: f"SUM({c}{P['med']}:{c}{P['cbfee']})", bold=True)
        P["tvc"] = r; r += 2
        self.month_row(ws, r, "Contribution profit", lambda c: f"{c}{P['net']}-{c}{P['tvc']}", bold=True)
        P["contrib"] = r; r += 1
        tcol = L(2 + self.nm)
        self.month_row(ws, r, "Contribution margin %", lambda c: f"IFERROR({c}{P['contrib']}/{c}{P['net']},0)", PCT,
                       total=f"IFERROR({tcol}{P['contrib']}/{tcol}{P['net']},0)")
        P["margin"] = r; r += 2
        self.month_row(ws, r, "Marketing spend", lambda c: f"SUM(Inputs!{c}${R['spend_r0']}:{c}${R['spend_r1']})")
        P["mkt"] = r; r += 1
        self.month_row(ws, r, "Fixed opex", lambda c: f"SUM(Inputs!{c}${R['opex_r0']}:{c}${R['opex_r1']})")
        P["opex"] = r; r += 2
        self.month_row(ws, r, "Operating profit", lambda c: f"{c}{P['contrib']}-{c}{P['mkt']}-{c}{P['opex']}", bold=True)
        P["op"] = r
        self.P = P

    # ------------------------------------------------------------------ Cash flow
    def cash_flow(self):
        ws, R, P = self.wb["Cash_Flow"], self.R, self.P
        Pay, S, F = self.D["Data_Payments"], self.D["Data_Shipments"], self.D["Data_Refunds"]
        self.month_header(ws, "Operating cash flow (cash timing, not accrual)")
        C, r = {}, 5
        pn = "Monthly_PnL"
        self.month_row(ws, r, "Card payouts (T+n settlement)",
                       lambda c: f"SUMIFS({Pay['amount']},{Pay['payout_idx']},{c}$4)"); C["payouts"] = r; r += 1
        self.month_row(ws, r, "Refunds paid", lambda c: f"{pn}!{c}{P['refunds']}"); C["refunds"] = r; r += 1
        self.month_row(ws, r, "Chargebacks debited", lambda c: f"{pn}!{c}{P['cbs']}"); C["cbs"] = r; r += 1
        self.month_row(ws, r, "Cash receipts", lambda c: f"{c}{C['payouts']}-{c}{C['refunds']}-{c}{C['cbs']}",
                       bold=True); C["receipts"] = r; r += 2
        section(ws, r, "Cash payments", 2 + self.nm); r += 1
        self.month_row(ws, r, "Pharmacy (medication + dispensing, net terms)",
                       lambda c: f"SUMIFS({S['pharmacy']},{S['paid_idx']},{c}$4)"); C["pharm"] = r; r += 1
        mk = f"{pn}!$B${P['mkt']}:${mcol(self.nm)}${P['mkt']}"
        self.month_row(ws, r, "Marketing (paid with lag)",
                       lambda c: f"IF({c}$4-{R['ad_lag']}>=1,INDEX({mk},1,{c}$4-{R['ad_lag']}),0)"); C["mkt"] = r; r += 1
        self.month_row(ws, r, "Other variable (consults, shipping, packaging, processing, CB fees)",
                       lambda c: (f"{pn}!{c}{P['consult']}+{pn}!{c}{P['ship']}+{pn}!{c}{P['pack']}"
                                  f"+{pn}!{c}{P['proc']}+{pn}!{c}{P['cbfee']}")); C["other"] = r; r += 1
        self.month_row(ws, r, "Fixed opex", lambda c: f"{pn}!{c}{P['opex']}"); C["opex"] = r; r += 1
        self.month_row(ws, r, "Cash payments", lambda c: f"SUM({c}{C['pharm']}:{c}{C['opex']})", bold=True)
        C["payments"] = r; r += 2
        self.month_row(ws, r, "Net operating cash flow", lambda c: f"{c}{C['receipts']}-{c}{C['payments']}",
                       bold=True); C["ncf"] = r; r += 2
        ws.cell(r, 1, "Opening cash")
        for m in range(1, self.nm + 1):
            col = mcol(m)
            v = R["opening_cash"] if m == 1 else f"{mcol(m - 1)}{r + 1}"
            ws.cell(r, 1 + m, f"={v}").number_format = USD
        C["open"] = r; r += 1
        ws.cell(r, 1, "Closing cash").font = BOLD
        for m in range(1, self.nm + 1):
            col = mcol(m)
            c = ws.cell(r, 1 + m, f"={col}{C['open']}+{col}{C['ncf']}")
            c.number_format, c.font = USD, BOLD
        C["close"] = r; r += 2
        self.month_row(ws, r, "Gap: cash flow minus operating profit",
                       lambda c: f"{c}{C['ncf']}-{pn}!{c}{P['op']}"); C["gap"] = r
        self.C = C

    # ------------------------------------------------------------------ Channels
    def channels_sheet(self):
        ws, R = self.wb["Channels"], self.R
        Cu = self.D["Data_Customers"]
        ws["A1"] = "Channels: spend, new customers, CPA vs target, payback, retention"
        ws["A1"].font = TITLE
        header(ws, 3, ["Channel", "Period spend", "New paying customers", "CPA", "Target CPA",
                       "CPA vs target", "Max CPA (payback window)", "Break-even CPA", "Payback (months)",
                       "LTV / CPA"])
        r = 4
        self.CH = {}
        for ch in self.channels:
            ws.cell(r, 1, ch)
            ws.cell(r, 2, f"=INDEX({R['ch_spend_total']},MATCH($A{r},{R['ch_names']},0))").number_format = USD
            ws.cell(r, 3, f"=COUNTIFS({Cu['channel']},$A{r})").number_format = INT
            ws.cell(r, 4, f"=IFERROR(B{r}/C{r},0)").number_format = USD2
            self.CH[ch] = r
            r += 1
        self.ch_last = r - 1
        ws.cell(r, 1, "PAID_ONLY").font = BOLD
        ws.cell(r, 2, f"=SUMIFS({R['ch_spend_total']},{R['ch_paid']},1)").number_format = USD
        paid_new = "+".join(f"IF(INDEX({R['ch_paid']},MATCH(A{self.CH[ch]},{R['ch_names']},0))=1,C{self.CH[ch]},0)"
                            for ch in self.channels)
        ws.cell(r, 3, f"={paid_new}").number_format = INT
        ws.cell(r, 4, f"=IFERROR(B{r}/C{r},0)").number_format = USD2
        self.CH["PAID_ONLY"] = r; r += 1
        ws.cell(r, 1, "ALL_BLENDED").font = BOLD
        ws.cell(r, 2, f"=SUM(B4:B{self.ch_last})").number_format = USD
        ws.cell(r, 3, f"=SUM(C4:C{self.ch_last})").number_format = INT
        ws.cell(r, 4, f"=IFERROR(B{r}/C{r},0)").number_format = USD2
        self.CH["ALL_BLENDED"] = r
        r += 2

        section(ws, r, "New paying customers by month", 11); r += 1
        header(ws, r, ["Channel"] + self.months)
        idx_r = r + 1
        ws.cell(idx_r, 1, "Month #").font = Font(italic=True, color="808080")
        for m in range(1, self.nm + 1):
            ws.cell(idx_r, 1 + m, f"=INDEX({R['month_idx_row']},1,{m})").font = Font(color="808080")
        r += 2
        self.newm0 = r
        for ch in self.channels:
            ws.cell(r, 1, ch)
            for m in range(1, self.nm + 1):
                col = mcol(m)
                ws.cell(r, 1 + m, f"=COUNTIFS({Cu['channel']},$A{r},{Cu['signup_idx']},{col}${idx_r})").number_format = INT
            r += 1
        self.newm1 = r - 1
        ws.cell(r, 1, "Total").font = BOLD
        for m in range(1, self.nm + 1):
            col = mcol(m)
            ws.cell(r, 1 + m, f"=SUM({col}{self.newm0}:{col}{self.newm1})").number_format = INT
        self.newm_total = r
        self.newm_idx_r = idx_r
        r += 1
        ws.cell(r, 1, "Blended CPA")
        for m in range(1, self.nm + 1):
            col = mcol(m)
            ws.cell(r, 1 + m, f"=IFERROR(Monthly_PnL!{col}{self.P['mkt']}/{col}{self.newm_total},0)").number_format = USD2
        r += 2

        section(ws, r, "Retention by channel (share of eligible customers active, months since signup)", 11); r += 1
        rc = self.D["Data_RetentionChannel"]
        nk = self.nm
        header(ws, r, ["Channel"] + list(range(nk)))
        kr = r
        r += 1
        for ch in self.channels:
            ws.cell(r, 1, ch)
            for k in range(nk):
                ws.cell(r, 2 + k, (f"=IFERROR(SUMIFS({rc['retained']},{rc['channel']},$A{r},{rc['months_since_signup']},"
                                   f"{L(2 + k)}${kr})/SUMIFS({rc['eligible_customers']},{rc['channel']},$A{r},"
                                   f"{rc['months_since_signup']},{L(2 + k)}${kr}),\"\")")).number_format = PCT
            r += 1
        widths(ws, 22, 14, 12)
        ws.freeze_panes = "B4"

    # ------------------------------------------------------------------ Calc_LTV
    def calc_ltv(self):
        ws, R = self.wb["Calc_LTV"], self.R
        Cu, rc, fit = self.D["Data_Customers"], self.D["Data_RetentionChannel"], self.D["Data_RetentionFit"]
        h = self.a["ltv"]["horizon_months"]
        ws["A1"] = "LTV projection blocks: one per scenario x channel (12-month horizon, discounted)"
        ws["A1"].font = TITLE
        ws["A2"] = ("r(k) = fitted retention curve (Data_RetentionFit) less the scenario's retention drop; "
                    "value(k) = r x paid share x (revenue - variable cost) per paid month x discount factor.")
        scen = [("base", "0", "0"), ("medication +15%", R["sens_med"], "0"),
                ("retention -10pp", "0", R["sens_ret"]), ("both", R["sens_med"], R["sens_ret"])]
        self.scen_names = [s[0] for s in scen]
        # summary table
        header(ws, 4, ["Scenario", "Channel", "LTV12 revenue", "LTV12 contribution", "Break-even CPA",
                       "Target CPA", "Max CPA", "Actual CPA", "Payback (months)", "CPA vs target"])
        self.LS = {}
        sr = 5
        block_r = 5 + len(scen) * len(self.channels) + 3
        cols = ["k", "Calendar month #", "First month (1/0)", "Fitted retention", "Retention (scenario)",
                "Paid share", "Monthly-plan charge", "Gross per paid month", "Revenue per paid month",
                "Medication", "Processing", "Consult", "Variable cost per paid month", "Discount factor",
                "Revenue value", "Contribution value", "Cumulative contribution"]
        start_idx = f"MATCH({R['ltv_start']},{R['months']},0)"
        for sname, uplift, drop in scen:
            for ch in self.channels:
                b = block_r
                ws.cell(b, 1, "Scenario").font = BOLD
                ws.cell(b, 2, sname)
                ws.cell(b, 3, "Channel").font = BOLD
                ws.cell(b, 4, ch)
                ws.cell(b, 5, "Med uplift"); ws.cell(b, 6, f"={uplift}").number_format = PCT
                ws.cell(b, 7, "Retention drop"); ws.cell(b, 8, f"={drop}").number_format = PCT
                chc = f"$D${b}"
                p = b + 1
                # parameter row (label/value pairs)
                pr = [("Prepaid mix", f"=COUNTIFS({Cu['channel']},{chc},{Cu['plan_type']},\"prepaid_3m\")/COUNTIFS({Cu['channel']},{chc})", PCT),
                      ("Promo share", f"=IFERROR(COUNTIFS({Cu['channel']},{chc},{Cu['plan_type']},\"monthly\",{Cu['promo']},1)/COUNTIFS({Cu['channel']},{chc},{Cu['plan_type']},\"monthly\"),0)", PCT),
                      ("Paid share", f"=SUMIFS({rc['retained_paid']},{rc['channel']},{chc},{rc['months_since_signup']},\">=1\")/SUMIFS({rc['retained']},{rc['channel']},{chc},{rc['months_since_signup']},\">=1\")", PCT),
                      ("Actual CPA", f"=INDEX(Channels!$D$4:$D${self.ch_last},MATCH({chc},Channels!$A$4:$A${self.ch_last},0))", USD2)]
                for j, (lab, f, fmt) in enumerate(pr):
                    ws.cell(p, 1 + 2 * j, lab)
                    ws.cell(p, 2 + 2 * j, f).number_format = fmt
                mix, promo, paid, cpa = f"$B${p}", f"$D${p}", f"$F${p}", f"$H${p}"
                header(ws, p + 1, cols)
                k0, k1 = p + 2, p + 1 + h
                for i in range(h):
                    rr = k0 + i
                    first = f"$C{rr}"
                    f = {
                        "A": f"{i}",
                        "B": f"=MIN({start_idx}+A{rr},{R['n_months']})",
                        "C": f"=IF(A{rr}=0,1,0)",
                        "D": f"=INDEX(Data_RetentionFit!$D$2:${L(3 + h)}${1 + len(self.channels)},"
                             f"MATCH({chc},{fit['channel']},0),A{rr}+1)",
                        "E": f"=IF(A{rr}=0,D{rr},MAX(0,D{rr}-$H${b}))",
                        "F": f"=IF(A{rr}=0,1,{paid})",
                        "G": f"={R['price_m']}*(1-{promo}*{R['promo_disc']}*{first})",
                        "H": f"=(1-{mix})*G{rr}+{mix}*{R['price_p']}/{R['term_p']}",
                        "I": (f"=H{rr}*(1-IF({first}=1,{R['ref_first']},{R['ref_renew']})*{R['ref_share']}"
                              f"-IF({first}=1,{R['cb_first']},{R['cb_renew']}))"),
                        "J": (f"=INDEX({R['med_base']},1,B{rr})*INDEX({R['dose']},1,MIN(A{rr}+1,MAX({R['dose_tiers']})))"
                              f"*(1+$F${b})"),
                        "K": (f"=(1-{mix})*(G{rr}*{R['proc_pct']}+{R['proc_fixed']})"
                              f"+{mix}*({R['price_p']}*{R['proc_pct']}+{R['proc_fixed']})/{R['term_p']}"),
                        "L": f"=IF({first}=1,{R['consult_alloc']},{R['consult_fu']})",
                        "M": (f"=J{rr}+{R['disp']}+L{rr}+{R['shipping']}+{R['packaging']}+K{rr}"
                              f"+IF({first}=1,{R['cb_first']},{R['cb_renew']})*{R['cb_fee']}"),
                        "N": f"=1/(1+{R['disc_rate']})^A{rr}",
                        "O": f"=E{rr}*F{rr}*I{rr}*N{rr}",
                        "P": f"=E{rr}*F{rr}*(I{rr}-M{rr})*N{rr}",
                        "Q": f"=SUM(P${k0}:P{rr})",
                    }
                    for col, v in f.items():
                        c = ws[f"{col}{rr}"]
                        c.value = int(v) if col == "A" else v
                        if col in "DEFN":
                            c.number_format = "0.0000"
                        elif col in "GHIJKLMOPQ":
                            c.number_format = USD2
                # results
                res = b + h + 3
                n_below = f"COUNTIFS($Q${k0}:$Q${k1},\"<\"&{cpa})"
                results = [
                    ("LTV12 revenue", f"=SUM(O{k0}:O{k1})", USD2),
                    ("LTV12 contribution", f"=SUM(P{k0}:P{k1})", USD2),
                    ("Target CPA", f"=D{res}/{R['target_ratio']}", USD2),
                    ("Max CPA", f"=INDEX($Q${k0}:$Q${k1},{R['max_payback']})", USD2),
                    ("Payback (months)", (f"=IF({n_below}>={h},\"n/a\",{n_below}+({cpa}-IF({n_below}=0,0,"
                                          f"INDEX($Q${k0}:$Q${k1},{n_below})))/INDEX($P${k0}:$P${k1},{n_below}+1))"), NUM2),
                ]
                for j, (lab, f, fmt) in enumerate(results):
                    ws.cell(res, 1 + 2 * j, lab).font = BOLD
                    c = ws.cell(res, 2 + 2 * j, f)
                    c.number_format, c.font = fmt, BOLD
                cells = {"mix": f"Calc_LTV!{mix}", "promo": f"Calc_LTV!{promo}", "paid": f"Calc_LTV!{paid}",
                         "ltv_rev": f"Calc_LTV!$B${res}", "ltv_con": f"Calc_LTV!$D${res}",
                         "target": f"Calc_LTV!$F${res}", "max": f"Calc_LTV!$H${res}",
                         "payback": f"Calc_LTV!$J${res}", "cpa": f"Calc_LTV!{cpa}"}
                self.LS[(sname, ch)] = cells
                # summary row
                ws.cell(sr, 1, sname); ws.cell(sr, 2, ch)
                for j, key in enumerate(["ltv_rev", "ltv_con", "ltv_con", "target", "max", "cpa", "payback"]):
                    c = ws.cell(sr, 3 + j, f"={cells[key]}")
                    c.number_format = NUM2 if key == "payback" else USD2
                ws.cell(sr, 10, f"=IFERROR(H{sr}/F{sr}-1,0)").number_format = PCT
                sr += 1
                block_r = res + 3
        widths(ws, 18, 15, 17)
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ Unit economics
    def unit_economics(self):
        ws, R, P = self.wb["Unit_Economics"], self.R, self.P
        O, F = self.D["Data_Orders"], self.D["Data_Refunds"]
        self.month_header(ws, "Unit economics: AOV, CPA, LTV, CPA guardrails, sensitivity", total=False)
        header(ws, 3, ["AOV by month"] + self.months)
        U, r = {}, 5
        rec = lambda c, first: (f"SUMIFS({O['net_pm']},{O['order_idx']},\"<=\"&{c}$4,{O['last_idx']},\">=\"&{c}$4,"
                                f"{O['is_first']},{first})-SUMIFS({F['amount']},{F['refund_idx']},{c}$4,{F['is_first']},{first})")
        cnt = lambda c, first: f"SUMIFS({O['orders']},{O['order_idx']},{c}$4,{O['is_first']},{first})"
        pn = "Monthly_PnL"
        rows = [
            ("orders", "Completed orders", lambda c: f"{pn}!{c}{P['orders']}", INT),
            ("net", "Net revenue", lambda c: f"{pn}!{c}{P['net']}", USD),
            ("aov", "AOV (net revenue / completed orders)", lambda c: f"IFERROR({c}{U['net']}/{c}{U['orders']},0)", USD2),
            ("f_orders", "First orders", lambda c: cnt(c, 1), INT),
            ("f_rev", "First-order net revenue", lambda c: rec(c, 1), USD),
            ("f_aov", "First-order AOV", lambda c: f"IFERROR({c}{U['f_rev']}/{c}{U['f_orders']},\"\")", USD2),
            ("r_orders", "Repeat orders", lambda c: cnt(c, 0), INT),
            ("r_rev", "Repeat-order net revenue", lambda c: rec(c, 0), USD),
            ("r_aov", "Repeat-order AOV", lambda c: f"IFERROR({c}{U['r_rev']}/{c}{U['r_orders']},\"\")", USD2),
            ("ticket", "Avg net charge per order (ticket)",
             lambda c: f"IFERROR(SUMIFS({O['net']},{O['order_idx']},{c}$4)/{c}{U['orders']},0)", USD2),
        ]
        for key, label, fn, fmt in rows:
            U[key] = r
            self.month_row(ws, r, label, fn, fmt, total=None)
            r += 1
        r += 1

        section(ws, r, "Acquisition, LTV and CPA guardrails by channel (12-month projection, base case)", 14); r += 1
        header(ws, r, ["Channel", "Spend", "New customers", "CPA", "First-order contribution",
                       "Observed revenue LTV", "Observed contribution LTV", "Projected LTV12 revenue",
                       "Projected LTV12 contribution", "Break-even CPA", "Target CPA", "Max CPA",
                       "Payback (months)", "LTV12 / CPA"])
        r += 1
        lt, fo = self.D["Data_LTV"], self.D["Data_FirstOrders"]
        self.UE = {}
        for ch in self.channels + ["PAID_ONLY", "ALL_BLENDED"]:
            cr = self.CH[ch]
            ws.cell(r, 1, ch)
            ws.cell(r, 2, f"=Channels!B{cr}").number_format = USD
            ws.cell(r, 3, f"=Channels!C{cr}").number_format = INT
            ws.cell(r, 4, f"=Channels!D{cr}").number_format = USD2
            if ch in self.channels:
                ls = self.LS[("base", ch)]
                ws.cell(r, 5, (f"=(SUMIFS({fo['net_charges']},{fo['channel']},$A{r})-SUMIFS({fo['refunds']},{fo['channel']},$A{r})"
                               f"-SUMIFS({fo['order_var_cost']},{fo['channel']},$A{r})-SUMIFS({fo['cb_fees']},{fo['channel']},$A{r}))"
                               f"/SUMIFS({fo['first_orders']},{fo['channel']},$A{r})")).number_format = USD2
                ws.cell(r, 6, (f"=(SUMIFS({lt['net_charges']},{lt['channel']},$A{r})-SUMIFS({lt['refunds']},{lt['channel']},$A{r}))"
                               f"/SUMIFS({lt['customers']},{lt['channel']},$A{r})")).number_format = USD2
                ws.cell(r, 7, (f"=(SUMIFS({lt['net_charges']},{lt['channel']},$A{r})-SUMIFS({lt['refunds']},{lt['channel']},$A{r})"
                               f"-SUMIFS({lt['order_var_cost']},{lt['channel']},$A{r})-SUMIFS({lt['cb_fees']},{lt['channel']},$A{r}))"
                               f"/SUMIFS({lt['customers']},{lt['channel']},$A{r})")).number_format = USD2
                ws.cell(r, 8, f"={ls['ltv_rev']}").number_format = USD2
                ws.cell(r, 9, f"={ls['ltv_con']}").number_format = USD2
                ws.cell(r, 10, f"=I{r}").number_format = USD2
                ws.cell(r, 11, f"={ls['target']}").number_format = USD2
                ws.cell(r, 12, f"={ls['max']}").number_format = USD2
                ws.cell(r, 13, f"={ls['payback']}").number_format = NUM2
                ws.cell(r, 14, f"=IFERROR(I{r}/D{r},0)").number_format = NUM2
                # write back to Channels sheet
                chs = self.wb["Channels"]
                chs.cell(cr, 5, f"=Unit_Economics!K{r}").number_format = USD2
                chs.cell(cr, 6, f"=IFERROR(D{cr}/E{cr}-1,0)").number_format = PCT
                chs.cell(cr, 7, f"=Unit_Economics!L{r}").number_format = USD2
                chs.cell(cr, 8, f"=Unit_Economics!J{r}").number_format = USD2
                chs.cell(cr, 9, f"=Unit_Economics!M{r}").number_format = NUM2
                chs.cell(cr, 10, f"=Unit_Economics!N{r}").number_format = NUM2
            self.UE[ch] = r
            r += 1
        r += 1

        section(ws, r, "Observed LTV to date by cohort (per acquired customer)", 14); r += 1
        header(ws, r, ["Cohort", "Customers", "Months observed", "Revenue LTV", "Contribution LTV"])
        r += 1
        for m, label in enumerate(self.months, start=1):
            ws.cell(r, 1, label)
            ws.cell(r, 2, f"=SUMIFS({lt['customers']},{lt['cohort_idx']},{m})").number_format = INT
            ws.cell(r, 3, f"={R['n_months']}-{m}+1")
            ws.cell(r, 4, (f"=IFERROR((SUMIFS({lt['net_charges']},{lt['cohort_idx']},{m})-SUMIFS({lt['refunds']},{lt['cohort_idx']},{m}))"
                           f"/B{r},0)")).number_format = USD2
            ws.cell(r, 5, (f"=IFERROR((SUMIFS({lt['net_charges']},{lt['cohort_idx']},{m})-SUMIFS({lt['refunds']},{lt['cohort_idx']},{m})"
                           f"-SUMIFS({lt['order_var_cost']},{lt['cohort_idx']},{m})-SUMIFS({lt['cb_fees']},{lt['cohort_idx']},{m}))/B{r},0)")).number_format = USD2
            r += 1
        r += 1

        section(ws, r, "Sensitivity (12-month projection)", 14); r += 1
        metrics = [("LTV12 contribution", "ltv_con", USD2), ("Target CPA", "target", USD2),
                   ("Payback (months)", "payback", NUM2)]
        for title, key, fmt in metrics:
            header(ws, r, [title] + self.scen_names)
            r += 1
            for ch in self.channels:
                ws.cell(r, 1, ch)
                for j, s in enumerate(self.scen_names):
                    ws.cell(r, 2 + j, f"={self.LS[(s, ch)][key]}").number_format = fmt
                r += 1
            r += 1
        widths(ws, 36, 15, 15)

    # ------------------------------------------------------------------ Cohorts
    def cohorts(self):
        ws, R = self.wb["Cohorts"], self.R
        Cu, co, rn = self.D["Data_Customers"], self.D["Data_Cohorts"], self.D["Data_Renewal"]
        ws["A1"] = "Cohort retention: share active (payment covers month or active pause); blank = not yet observable"
        ws["A1"].font = TITLE
        header(ws, 3, ["Cohort", "Cohort #", "Customers"] + [f"M{k}" for k in range(self.nm)])
        ws.cell(4, 1, "Months since signup").font = Font(italic=True, color="808080")
        for k in range(self.nm):
            ws.cell(4, 4 + k, k).font = Font(color="808080")
        r0 = 5
        for i, label in enumerate(self.months):
            r = r0 + i
            ws.cell(r, 1, label)
            ws.cell(r, 2, f"=INDEX({R['month_idx_row']},1,{i + 1})")
            ws.cell(r, 3, f"=COUNTIFS({Cu['signup_idx']},B{r})").number_format = INT
            for k in range(self.nm):
                col = L(4 + k)
                ws.cell(r, 4 + k, (f"=IF($B{r}+{col}$4>{R['n_months']},\"\",SUMIFS({co['retained']},{co['cohort_idx']},"
                                   f"$B{r},{co['months_since_signup']},{col}$4)/$C{r})")).number_format = PCT
        r1 = r0 + self.nm - 1
        self.cohort_rows = (r0, r1)
        ws.conditional_formatting.add(f"D{r0}:{L(3 + self.nm)}{r1}",
                                      ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                     mid_type="num", mid_value=0.5, mid_color="FFEB84",
                                                     end_type="num", end_value=1, end_color="63BE7B"))
        r = r1 + 2
        section(ws, r, "Renewal rate and monthly churn", 12); r += 1
        header(ws, r, ["Month", "Month #", "Active customers", "Active prev. month", "Churned", "Monthly churn",
                       "Renewals due", "Renewed", "Renewal rate"])
        r += 1
        for i, label in enumerate(self.months):
            ws.cell(r, 1, label)
            ws.cell(r, 2, f"=INDEX({R['month_idx_row']},1,{i + 1})")
            for j, col in enumerate(["active_customers", "active_prev_month", "churned"]):
                ws.cell(r, 3 + j, f"=SUMIFS({rn[col]},{rn['month_idx']},$B{r})").number_format = INT
            ws.cell(r, 6, f"=IFERROR(E{r}/D{r},\"\")").number_format = PCT
            for j, col in enumerate(["renewals_due", "renewed"]):
                ws.cell(r, 7 + j, f"=SUMIFS({rn[col]},{rn['month_idx']},$B{r})").number_format = INT
            ws.cell(r, 9, f"=IFERROR(H{r}/G{r},\"\")").number_format = PCT
            r += 1
        widths(ws, 20, 11, 14)
        ws.freeze_panes = "D5"

    # ------------------------------------------------------------------ Checks
    def checks(self):
        ws, R, P, C = self.wb["Checks"], self.R, self.P, self.C
        S, O, Pay, Cu = (self.D["Data_Shipments"], self.D["Data_Orders"], self.D["Data_Payments"],
                         self.D["Data_Customers"])
        ws["A1"] = "Checks: reconciliations (TRUE = passes within tolerance)"
        ws["A1"].font = TITLE
        header(ws, 4, ["Check", "Value A", "Value B", "Difference", "Pass"])
        rows = []
        cf, pn, tc = "Cash_Flow", "Monthly_PnL", L(2 + self.nm)
        for m in range(1, self.nm + 1):
            c = mcol(m)
            rows.append((f"Cash roll-forward {self.months[m - 1]}: opening + net flow = closing",
                         f"{cf}!{c}{C['open']}+{cf}!{c}{C['ncf']}", f"{cf}!{c}{C['close']}"))
        for m in range(1, self.nm):
            rows.append((f"Closing cash {self.months[m - 1]} = opening cash {self.months[m]}",
                         f"{cf}!{mcol(m)}{C['close']}", f"{cf}!{mcol(m + 1)}{C['open']}"))
        rows.append(("Opening cash = input", f"{cf}!B{C['open']}", R["opening_cash"]))
        for m in range(1, self.nm + 1):
            c = mcol(m)
            rows.append((f"New customers by channel sum to total {self.months[m - 1]}",
                         f"Channels!{c}{self.newm_total}", f"COUNTIFS({Cu['signup_idx']},{m})"))
        rows.append(("Channel table new customers = all customers",
                     f"Channels!C{self.CH['ALL_BLENDED']}", f"COUNTIFS({Cu['signup_idx']},\">=1\")"))
        rows.append(("Channel spend total = P&L marketing", f"Channels!B{self.CH['ALL_BLENDED']}", f"{pn}!{tc}{P['mkt']}"))
        rows.append(("Order revenue at input prices = billed amounts in data", f"SUM({O['net']})", f"SUM({O['net_billed_data']})"))
        rows.append(("Payments at input prices = payment amounts in data", f"SUM({Pay['amount']})", f"SUM({Pay['amount_data']})"))
        rows.append(("P&L net revenue = gross - discounts - refunds - chargebacks", f"{pn}!{tc}{P['net']}",
                     f"{pn}!{tc}{P['gross']}-{pn}!{tc}{P['disc']}-{pn}!{tc}{P['refunds']}-{pn}!{tc}{P['cbs']}"))
        rows.append(("Total variable cost = sum of cost lines", f"{pn}!{tc}{P['tvc']}", f"SUM({pn}!{tc}{P['med']}:{tc}{P['cbfee']})"))
        rows.append(("P&L medication = shipment medication (counted once)", f"{pn}!{tc}{P['med']}", f"SUM({S['medication']})"))
        rows.append(("Contribution = net revenue - variable cost", f"{pn}!{tc}{P['contrib']}", f"{pn}!{tc}{P['net']}-{pn}!{tc}{P['tvc']}"))
        rows.append(("Operating profit = contribution - marketing - opex", f"{pn}!{tc}{P['op']}",
                     f"{pn}!{tc}{P['contrib']}-{pn}!{tc}{P['mkt']}-{pn}!{tc}{P['opex']}"))
        rows.append(("Pharmacy accrued = paid in period + payable at period end",
                     f"{pn}!{tc}{P['med']}+{pn}!{tc}{P['disp']}",
                     f"{cf}!{tc}{C['pharm']}+SUMIFS({S['pharmacy']},{S['paid_idx']},\">\"&{R['n_months']})"))
        r0, r1 = self.cohort_rows
        rows.append(("Cohort sizes sum to all customers", f"SUM(Cohorts!C{r0}:C{r1})", f"COUNTIFS({Cu['signup_idx']},\">=1\")"))
        rows.append(("Cohort month-0 retention = 100% (min)", f"MIN(Cohorts!D{r0}:D{r1})", "1"))
        rows.append(("Cohort month-0 retention = 100% (max)", f"MAX(Cohorts!D{r0}:D{r1})", "1"))
        fs = self.fc.S
        rows.append(("Forecast opening cash = last actual closing cash", f"Forecast!$D${fs['opening']}",
                     f"INDEX(Cash_Flow!$B${C['close']}:${mcol(self.nm)}${C['close']},1,{R['n_months']})"))
        rows.append(("Forecast month-1 active start = actual active customers at period end",
                     f"SUM({self.D['Data_ActiveEnd']['active_customers']})",
                     f"SUMIFS({self.D['Data_Renewal']['active_customers']},{self.D['Data_Renewal']['month_idx']},{R['n_months']})"))
        rows.append(("Bridge bars sum to total change", f"Doubling_Bridge!$B${self.fc.bridge_total_row}",
                     "Doubling_Bridge!$B$9"))
        rows.append(("Forecast = scenario table for the selected scenario (last month cash flow)",
                     f"INDEX(Forecast!$D${fs['ncf']}:${L(3 + NF)}${fs['ncf']},1,{NF})",
                     f"IF({R['scenario']}=1,INDEX(Eng_Base!$D${fs['ncf']}:${L(3 + NF)}${fs['ncf']},1,{NF}),"
                     f"INDEX(Eng_Down!$D${fs['ncf']}:${L(3 + NF)}${fs['ncf']},1,{NF}))"))
        rows.append(("No cohort cell above 100%", f"MAX(Cohorts!D{r0}:{L(3 + self.nm)}{r1})", "MIN(1,MAX(Cohorts!D"
                     f"{r0}:{L(3 + self.nm)}{r1}))"))
        r = 5
        for label, fa, fb in rows:
            ws.cell(r, 1, label)
            ws.cell(r, 2, f"={fa}").number_format = USD2
            ws.cell(r, 3, f"={fb}").number_format = USD2
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = USD2
            ws.cell(r, 5, f"=MAX(D{r},-D{r})<={R['tol']}")
            ws.cell(r, 6, f"=IFERROR(IF(E{r}=TRUE,1,0),0)")
            r += 1
        ws["A3"] = "All checks pass (errors count as failures)"
        ws["A3"].font = BOLD
        ws["B3"] = f"=SUM(F5:F{r - 1})=COUNTIFS(A5:A{r - 1},\"*\")"
        ws["B3"].font = BOLD
        widths(ws, 62, 18, 5)
        ws.freeze_panes = "A5"

    # ------------------------------------------------------------------ README
    def readme(self):
        ws = self.wb["README"]
        lines = [
            ("Omzo GLP-1 telehealth: financial model (synthetic data, Jan-Sep 2026)", TITLE),
            ("Purpose: monthly P&L, operating cash flow, unit economics and cohort retention for a hypothetical "
             "US telehealth GLP-1 subscription business, as the base for a plan to double monthly net operating cash flow.", None),
            ("", None),
            ("Colour legend", BOLD),
            ("Blue font = input (edit on Inputs). Black font = live formula. Grey = structural labels (month numbers).", None),
            ("", None),
            ("How it works", BOLD),
            ("Data_* sheets hold volumes aggregated from data/*.csv (generated by src/generate_data.py). Prices, cost rates and "
             "timing rules on Inputs are applied by live formulas, so editing an input re-prices the P&L, cash flow and unit economics.", None),
            ("Inputs marked 'data' shaped the simulated volumes (churn, pause, refund rates...). Changing them here does not "
             "re-simulate; edit assumptions.yaml and re-run the pipeline instead.", None),
            ("Retention aggregates (Data_Cohorts, Data_Renewal, Data_RetentionChannel) come from the DuckDB SQL layer; "
             "fitted retention curves (Data_RetentionFit) come from src/ltv_projection.py.", None),
            ("Functions used: SUM, SUMIFS, COUNTIFS, INDEX, MATCH, IF, IFERROR, ROUND, MIN, MAX (Excel and Google Sheets).", None),
            ("", None),
            ("Sheet guide", BOLD),
            ("Inputs: every assumption from assumptions.yaml, with units and what it drives.", None),
            ("Monthly_PnL: net revenue, each variable cost, contribution and margin, marketing, fixed opex, operating profit.", None),
            ("Cash_Flow: receipts, payments by type, net operating cash flow, opening and closing cash, gap to profit.", None),
            ("Unit_Economics: AOV (first vs repeat), CPA, observed and projected LTV, CPA guardrails, payback, sensitivity.", None),
            ("Channels: spend, new customers, CPA vs target, payback, retention curve by channel.", None),
            ("Cohorts: retention triangle (colour scale), renewal rate and monthly churn.", None),
            ("Calc_LTV: the 12-month LTV projection, one block per scenario x channel.", None),
            ("Checks: reconciliations returning TRUE/FALSE; cell B3 summarises all.", None),
            ("Data_*: Calendar, Customers (row-level), Orders, Shipments, Payments, Refunds, LTV, FirstOrders, Cohorts, "
             "Renewal, RetentionChannel, RetentionFit.", None),
            ("", None),
            ("Metric definitions", BOLD),
            ("Net revenue: gross charges - promo discounts - refunds - chargebacks. Prepaid recognized ratably over 3 months.", None),
            ("AOV: net revenue / completed orders in the month (first/repeat split net revenue by originating order).", None),
            ("Channel CPA: channel marketing spend / new paying customers from that channel.", None),
            ("Blended CPA: total marketing spend (organic included) / total new paying customers.", None),
            ("Contribution: net revenue - all variable costs (medication, dispensing, consults, shipping, packaging, "
             "processing, chargeback fees). Excludes marketing and fixed opex.", None),
            ("Operating profit: contribution - marketing spend - fixed opex.", None),
            ("Net operating cash flow: operating cash receipts - operating cash payments in the month, after timing "
             "(card payouts T+2, pharmacy net-30, ads paid next month, refund/chargeback lags). Excludes financing and capex.", None),
            ("Retention: share of a signup cohort still active (payment covers the month or active pause) N months after signup.", None),
            ("Churn: share of customers active last month who are not active this month (voluntary cancel or failed payment).", None),
            ("LTV: expected cumulative contribution per new customer over the stated horizon (12 months here), discounted, before CPA.", None),
            ("", None),
            ("Cost categories (each cost in exactly one bucket)", BOLD),
            ("COGS: medication only (base cost of the ship month x dose for tenure). Dispensing excludes the drug.", None),
            ("Variable: fulfillment (dispensing, shipping, packaging), clinical (consults), payments (processing, chargeback fees).", None),
            ("Contra-revenue: promo discounts, refunds, chargeback amounts. Marketing and fixed opex sit below contribution.", None),
        ]
        for i, (text, font) in enumerate(lines, start=1):
            c = ws.cell(i, 1, text)
            if font:
                c.font = font
            c.alignment = Alignment(wrap_text=True, vertical="top")
        self.readme_plan(ws, len(lines) + 2)
        ws.column_dimensions["A"].width = 140
        for col in "BCDE":
            ws.column_dimensions[col].width = 16

    def readme_plan(self, ws, r):
        """Baseline choice, plan results (live), credibility check and concentration flags."""
        R, P, C, fc = self.R, self.P, self.C, self.fc
        wrap = Alignment(wrap_text=True, vertical="top")

        def text(t, font=None):
            nonlocal r
            c = ws.cell(r, 1, t)
            c.alignment = wrap
            if font:
                c.font = font
            r += 1

        text("Plan to double monthly net operating cash flow", TITLE)
        text("1. Baseline month: candidates = months with positive net operating cash flow", BOLD)
        header(ws, r, ["Month", "Net op. cash flow", "Operating profit", "Cash minus profit", "Positive?"])
        r += 1
        for m in range(1, self.nm + 1):
            col = mcol(m)
            ws.cell(r, 1, f"=INDEX({R['months']},1,{m})")
            ws.cell(r, 2, f"=Cash_Flow!{col}{C['ncf']}").number_format = USD
            ws.cell(r, 3, f"=Monthly_PnL!{col}{P['op']}").number_format = USD
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = USD
            ws.cell(r, 5, f"=B{r}>0")
            r += 1
        text("Choice: September 2026. Why not the others:")
        text("- January is an artefact. It was the first month, so no ad invoices were paid yet (ads are paid a month "
             "later), while the business was still loss-making.")
        text("- March, May and July are near zero. The spend ramp and June's peak-price pharmacy bills (paid in July) "
             "hit cash.")
        text("- August and September are consistent with each other (September is about 17% above August, in line "
             "with active-customer growth). Both reflect the post-July medication contract on accrual and cash. "
             "September is the latest full month and operating profit is positive, so it is representative, not a spike.")
        text("Caveat: about half of September's cash flow is working-capital timing from growth (prepaid cash ahead "
             "of revenue, pharmacy paid net-30). The forecast models that timing explicitly, so the target is "
             "compared like for like.")
        ws.cell(r, 1, "Baseline net operating cash flow (live)")
        ws.cell(r, 2, f"={fc.base_ncf}").number_format = USD
        r += 1
        ws.cell(r, 1, "Target = multiple x baseline, to reach within Oct 2026 - Jun 2027")
        ws.cell(r, 2, f"={fc.target}").number_format = USD
        r += 2

        text("2. Results (live; base case uses Levers as set, downside per Scenarios)", BOLD)
        sr = fc.sum_rows
        for lab in ["Month 2x first reached and sustained", "Peak additional cash needed (opening - lowest closing)",
                    "Net operating cash flow, last forecast month", "Gap to target, last forecast month"]:
            ws.cell(r, 1, lab)
            for j, col in enumerate("BC"):
                ws.cell(r, 2 + j, f"=Scenarios!{col}{sr[lab]}").number_format = USD
            r += 1
        ws.cell(r - 4, 4, "(base | downside)")
        r += 1
        text("3. Credibility check", BOLD)
        down = f"Scenarios!C{sr['Month 2x first reached and sustained']}"
        gap = f"Scenarios!C{sr['Gap to target, last forecast month']}"
        last = f"Scenarios!C{sr['Net operating cash flow, last forecast month']}"
        ws.cell(r, 1, (f"=IF({down}=\"not reached\",\"DOWNSIDE DOES NOT REACH 2X: last-month cash flow is $\"&ROUND({last},0)"
                       f"&\" vs target $\"&ROUND({fc.target},0)&\" (short by $\"&ROUND({gap},0)&\"/month).\","
                       f"\"Downside reaches 2x in \"&{down}&\".\")")).font = BOLD
        r += 1
        text("What would need to change (scripted what-ifs at default inputs, downside Jun-2027 cash flow): "
             "levers at 100% instead of 50%: ~$1.52M; no retention drop: ~$1.48M; no medication shock: ~$1.34M; "
             "no CPA shock: ~$1.29M. None reaches 2x alone. Only levers at full effect AND retention held at "
             "base levels (the three retention levers fully offsetting a 10pp deterioration) gets there: ~$2.21M, "
             "sustained from Apr 2027. So the plan depends on retention execution: it must be proven early with "
             "the check-in and dunning tests.")
        text("Plan-mix lever (c) is modelled but OFF by default. Switched on in the base case it costs ~$129k/month of "
             "cash (-$239k contribution) at the target month. Prepaid earns $284-299/month vs $349, and its "
             "retention edge does not offset that within the horizon. Recommendation: no extra prepaid discount; "
             "test prepaid at $897 only for cash timing.")
        r += 1
        text("4. Concentration risk (any lever > 40% of the uplift)", BOLD)
        flags = "&".join(f"IF(Doubling_Bridge!D{br}>0.4,Doubling_Bridge!A{br}&\" (\"&ROUND(Doubling_Bridge!D{br}*100,0)"
                         f"&\"%) \",\"\")" for br in fc.bar_rows[1:])
        ws.cell(r, 1, f"=IF(({flags})=\"\",\"No single lever exceeds 40% of the uplift.\",\"CONCENTRATION RISK: \"&{flags})")
        r += 1
        text("Momentum (growth of the active base at constant spend) is shown separately in the bridge; it is not a lever.")
        r += 1
        text("Forecast simplifications", BOLD)
        text("Refunds and chargebacks are booked in the month of the charge. Fixed opex is held flat (growth input on Levers). "
             "Medication base cost is held at the last actual price unless a lever or shock applies. Dose tier "
             "follows months since signup. Pharmacy is paid the following month. Pauses come only from lever b3; "
             "the data's existing pause behaviour is inside the paid-share rate.")


def build(out: Path = OUT, scenario: int = 1) -> Path:
    a = load_assumptions()
    wb = Model(a, scenario).build()
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def summary(path: Path) -> None:
    """Recalculate with LibreOffice and print the plan summary."""
    from recalc import find_soffice, recalc
    if not find_soffice():
        print("LibreOffice not installed: skipping recalculated summary")
        return
    v = recalc(path, path.parent / ".recalc")
    sc, b = v["Scenarios"], v["Doubling_Bridge"]
    get = lambda label, col: next(sc.cell(r, col).value for r in range(1, sc.max_row + 1)
                                  if sc.cell(r, 1).value == label)
    usd = lambda x: f"${x:,.0f}" if isinstance(x, (int, float)) else str(x)
    print()
    print(f"Baseline ({get('Baseline month', 2)}) net operating cash flow: {usd(get('Baseline net operating cash flow', 2))}")
    print(f"Target (2x): {usd(get('Target (multiple x baseline)', 2))}")
    for col, name in [(2, "Base"), (3, "Downside")]:
        print(f"{name:<9} 2x reached & sustained: {get('Month 2x first reached and sustained', col):<12} "
              f"Jun-27 cash flow: {usd(get('Net operating cash flow, last forecast month', col)):>12}  "
              f"peak extra cash needed: {usd(get('Peak additional cash needed (opening - lowest closing)', col))}")
    start = next(r for r in range(1, b.max_row + 1) if b.cell(r, 1).value == "Momentum (no levers)")
    end = next(r for r in range(1, b.max_row + 1) if b.cell(r, 1).value == "Target-month cash flow")
    bars = [(b.cell(r, 1).value, b.cell(r, 2).value, b.cell(r, 4).value) for r in range(start, end)]
    print(f"Bridge to {b['B4'].value}: momentum {usd(bars[0][1])} ({bars[0][2]:.0%} of uplift)")
    print("Top 3 levers by cash impact:")
    for name, val, pct in sorted(bars[1:], key=lambda x: -x[1])[:3]:
        print(f"  {name:<52}{usd(val):>12}/month  {pct:.0%} of uplift")
    flagged = [n for n, _, pct in bars[1:] if pct > 0.4]
    print("Concentration risk (>40%):", ", ".join(flagged) if flagged else "none")
    print("Workbook checks all pass:", v["Checks"]["B3"].value)
    export_forecast(v)


def export_forecast(v) -> None:
    """Write recalculated forecast results to output/metrics/ for charts and reports."""
    METRICS.mkdir(parents=True, exist_ok=True)
    rows = []
    labels = {"New paying customers": "new_customers", "Active customers (incl. paused)": "active_customers",
              "Net revenue": "net_revenue", "Contribution profit": "contribution", "Operating profit": "operating_profit",
              "Net operating cash flow": "net_operating_cash_flow", "Closing cash": "closing_cash"}
    for sheet, scen in [("Eng_Base", "base"), ("Eng_Down", "downside"), ("Eng_Bridge0", "no_levers")]:
        ws = v[sheet]
        months = [ws.cell(19, 3 + f).value for f in range(1, NF + 1)]
        found = {labels[ws.cell(r, 1).value]: r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value in labels}
        for f, mth in enumerate(months, start=1):
            rows.append({"scenario": scen, "month": mth, **{k: ws.cell(r, 3 + f).value for k, r in found.items()}})
    pd.DataFrame(rows).to_csv(METRICS / "forecast_scenarios.csv", index=False)

    b, sc = v["Doubling_Bridge"], v["Scenarios"]
    start = next(r for r in range(1, b.max_row + 1) if b.cell(r, 1).value == "Momentum (no levers)")
    end = next(r for r in range(1, b.max_row + 1) if b.cell(r, 1).value == "Target-month cash flow")
    pd.DataFrame([{"step": b.cell(r, 1).value, "cash_impact": b.cell(r, 2).value,
                   "contribution_impact": b.cell(r, 3).value, "pct_of_uplift": b.cell(r, 4).value,
                   "timing": b.cell(r, 5).value, "confidence": b.cell(r, 6).value}
                  for r in range(start, end)]).to_csv(METRICS / "doubling_bridge.csv", index=False)
    get = lambda label, col: next(sc.cell(r, col).value for r in range(1, sc.max_row + 1)
                                  if sc.cell(r, 1).value == label)
    summ = {"baseline_month": get("Baseline month", 2),
            "baseline_ncf": get("Baseline net operating cash flow", 2),
            "target_ncf": get("Target (multiple x baseline)", 2),
            "bridge_month": b["B4"].value, "target_month_ncf": b["B9"].value}
    for col, scen in [(2, "base"), (3, "downside")]:
        summ[f"{scen}_month_reached"] = get("Month 2x first reached and sustained", col)
        summ[f"{scen}_peak_cash_need"] = get("Peak additional cash needed (opening - lowest closing)", col)
        summ[f"{scen}_last_month_ncf"] = get("Net operating cash flow, last forecast month", col)
    pd.DataFrame([summ]).to_csv(METRICS / "forecast_summary.csv", index=False)


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--out", type=Path, default=OUT)
    ap.add_argument("--no-summary", action="store_true", help="skip LibreOffice recalculation + summary")
    args = ap.parse_args()
    path = build(args.out)
    print(f"Saved {path}")
    if not args.no_summary:
        summary(path)
