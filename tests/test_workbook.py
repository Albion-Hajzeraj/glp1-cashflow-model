"""Workbook verification: rebuild output/omzo_model.xlsx, recalculate it with LibreOffice
headless (a real spreadsheet engine; openpyxl stores formulas without values), then
compare key outputs with the DuckDB/Python metrics in output/metrics/. Tolerance: < $1.
"""
import numpy as np
import pytest
from openpyxl import load_workbook

from conftest import ROOT

TOL = 1.0


@pytest.fixture(scope="module")
def wb(recalculated):
    return recalculated


ALLOWED_FUNCTIONS = {"SUM", "SUMIFS", "COUNTIFS", "AVERAGEIFS", "INDEX", "MATCH", "IF", "IFERROR",
                     "ROUND", "MIN", "MAX"}


def test_only_portable_functions_and_no_external_links(tmp_path):
    """Every formula uses only functions that work in both Excel and Google Sheets."""
    import re
    import build_workbook
    path = build_workbook.build(tmp_path / "omzo_model.xlsx")
    book = load_workbook(path)
    used, bad = set(), []
    for ws in book.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    f = re.sub(r'"[^"]*"', "", c.value)            # drop string literals
                    names = set(re.findall(r"([A-Z][A-Z0-9.]*)\(", f))
                    used |= names
                    if names - ALLOWED_FUNCTIONS or "[" in f or "{" in f:
                        bad.append(f"{ws.title}!{c.coordinate}: {c.value[:80]}")
    assert not bad, bad[:10]
    assert not book.defined_names, "named ranges not expected"
    assert used <= ALLOWED_FUNCTIONS


def row_by_label(ws, label, col=1):
    for r in range(1, ws.max_row + 1):
        if ws.cell(r, col).value == label:
            return r
    raise KeyError(f"{label!r} not found on {ws.title}")


def month_values(ws, label, n):
    r = row_by_label(ws, label)
    return np.array([ws.cell(r, 1 + m).value for m in range(1, n + 1)], dtype=float)


def assert_close(actual, expected, what):
    diff = np.abs(np.asarray(actual, dtype=float) - np.asarray(expected, dtype=float))
    assert (diff < TOL).all(), f"{what}: max diff ${diff.max():,.4f}"


def test_all_workbook_checks_pass(wb):
    ws = wb["Checks"]
    assert ws["B3"].value is True, [ws.cell(r, 1).value for r in range(5, ws.max_row + 1)
                                    if ws.cell(r, 5).value is not True]


def test_no_formula_errors(wb):
    errors = []
    for ws in wb.worksheets:
        if ws.title.startswith("Data_Customers"):
            continue
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#") and c.value.endswith(("!", "?", "A")):
                    errors.append(f"{ws.title}!{c.coordinate}={c.value}")
    assert not errors, errors[:20]


def test_monthly_pnl_matches_metrics(wb, metrics):
    pnl = metrics["monthly_pnl"]
    n = len(pnl)
    ws = wb["Monthly_PnL"]
    assert_close(month_values(ws, "Net revenue", n), pnl.net_revenue, "net revenue")
    assert_close(month_values(ws, "Contribution profit", n), pnl.contribution_profit, "contribution")
    assert_close(month_values(ws, "Total variable cost", n), pnl.total_variable_cost, "variable cost")
    assert_close(month_values(ws, "Operating profit", n), pnl.operating_profit, "operating profit")


def test_cash_flow_matches_metrics(wb, metrics):
    cf = metrics["cash_flow"]
    n = len(cf)
    ws = wb["Cash_Flow"]
    assert_close(month_values(ws, "Net operating cash flow", n), cf.net_operating_cash_flow, "net op cash flow")
    assert_close(month_values(ws, "Closing cash", n), cf.closing_cash, "closing cash")
    assert_close(month_values(ws, "Cash receipts", n), cf.cash_receipts, "cash receipts")


def test_blended_and_channel_cpa_match_metrics(wb, metrics):
    ws = wb["Channels"]
    per = metrics["acquisition_period"].set_index("channel")
    for ch in per.index:
        r = row_by_label(ws, ch)
        assert abs(ws.cell(r, 4).value - per.loc[ch, "cpa"]) < TOL, ch


def test_ltv_matches_metrics(wb, metrics):
    ws = wb["Unit_Economics"]
    proj = metrics["ltv_and_cpa"].set_index("channel")
    obs = metrics["ltv_observed"].query("cohort == 'ALL'").set_index("channel")
    foc = metrics["first_order_contribution"].set_index("channel")
    hdr = row_by_label(ws, "Channel")
    col = {ws.cell(hdr, j).value: j for j in range(1, 16)}
    for ch in proj.index:
        r = row_by_label(ws, ch)
        v = lambda name: ws.cell(r, col[name]).value
        assert abs(v("Projected LTV12 revenue") - proj.loc[ch, "ltv12_revenue"]) < TOL, ch
        assert abs(v("Projected LTV12 contribution") - proj.loc[ch, "ltv12_contribution"]) < TOL, ch
        assert abs(v("Target CPA") - proj.loc[ch, "target_cpa"]) < TOL, ch
        assert abs(v("Max CPA") - proj.loc[ch, "max_cpa_4m"]) < TOL, ch
        assert abs(v("Payback (months)") - proj.loc[ch, "payback_months"]) < 0.01, ch
        assert abs(v("Observed revenue LTV") - obs.loc[ch, "revenue_ltv_to_date"]) < TOL, ch
        assert abs(v("Observed contribution LTV") - obs.loc[ch, "contribution_ltv_to_date"]) < TOL, ch
        assert abs(v("First-order contribution") - foc.loc[ch, "first_order_contribution"]) < TOL, ch


def test_sensitivity_matches_metrics(wb, metrics):
    ws = wb["Calc_LTV"]
    sens = metrics["sensitivity"]
    names = {"base": "base", "medication +15%": "medication +15%",
             "retention -10pp": "retention -10pp", "both": "both"}
    rows = {(ws.cell(r, 1).value, ws.cell(r, 2).value): r for r in range(5, 5 + len(sens))}
    for s in sens.itertuples():
        r = rows[(names[s.scenario], s.channel)]
        assert abs(ws.cell(r, 4).value - s.ltv12_contribution) < TOL, (s.scenario, s.channel)
        assert abs(ws.cell(r, 6).value - s.target_cpa) < TOL, (s.scenario, s.channel)


def test_retention_triangle_matches_metrics(wb, metrics):
    ws = wb["Cohorts"]
    long = metrics["cohort_retention_long"]
    for t in long.itertuples():
        r = row_by_label(ws, t.cohort)
        v = ws.cell(r, 4 + t.months_since_signup).value
        if np.isnan(t.retention_pct):
            assert v in (None, ""), (t.cohort, t.months_since_signup, v)
        else:
            assert abs(v - t.retention_pct) < 1e-9, (t.cohort, t.months_since_signup)
